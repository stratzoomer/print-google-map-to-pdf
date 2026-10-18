#!/usr/bin/env python3
"""
analyze_deliveries.py
=====================

After the delivery day: who delivered what, and how fast, to help assign
routes to drivers next season.

The drivers are logged on the ``OrderDataByDeliveryRoute`` sheet, next to
each delivery route: ``Driver(s):``, ``Time Out`` (when the truck left the
Sideburn Run lot) and ``Time In`` (when it got back).  Routes a driver took
out together (same driver, same ``Time Out``) are one trip, which ends at
the latest ``Time In`` of its routes.  The stops and bags of each route are
counted from the orders on ``Customers`` (like ``check_workbook.py``).

First it checks the driver log: every route with bags must have a
``Time Out`` and a ``Time In``, and the ``Time In`` must be later than the
``Time Out``.  If not, it lists those routes and stops (exit 1) without a
report; fix the workbook and run again (``--skip-checks`` reports anyway,
guessing around the mistakes).

The report lists:

- problems in the driver log (missing times, ``Time In`` before ``Time
  Out``, routes with orders but no driver); fix them in the workbook and
  run again,
- per driver: stops, bags, minutes out, bags and stops per minute, the
  biggest load, and the pace compared with what's typical for trips of
  that size (a fitted ``minutes = a + b*stops + c*bags``),
- per trip: routes, stops, bags, miles (from ``stop_order.json``),
  minutes, expected minutes.

Per-minute numbers only count trips with both times; a driver's minutes
are the time they were out.  A route that leaves before the driver's trip
is back was added to that trip on the way (e.g. Fairfax 01A out 7:35, back
10:55, and Fairfax 02B out 9:47, back 10:55): they are one trip, from the
first ``Time Out`` to the last ``Time In``, but the biggest load is the
most bags that left at once.  A trip with several drivers counts fully for
each of them.

Usage::

    python analyze_deliveries.py --input "input/Mulch Sales - Fall 2026.xlsx"
    python analyze_deliveries.py --input "..." --stop-order output/stop_order.json \
        --csv-dir output/delivery_analysis
"""

import argparse
import csv
import datetime
import json
import os
import re
import sys
from typing import Dict, List, Optional, Sequence, Tuple

from check_workbook import ROUTE_STATS_SHEET, load_orders

# A route that leaves more than this many minutes before the driver's trip
# is back was added to that trip on the way (less: rounded times, a new trip).
OVERLAP_TOLERANCE_MIN = 5
# Times before this hour are taken as afternoon times typed without "pm".
EARLIEST_HOUR = 6
KM_PER_MILE = 1.609344


class Trip:
    """Routes one driver (or several together) delivered in one go.

    Usually they leave together.  A driver may also add a route to a trip
    on the way (it leaves before the trip is back): then the trip has
    several legs, one per ``Time Out``.
    """

    def __init__(self, drivers: Tuple[str, ...], time_out: Optional[int]):
        self.drivers = drivers
        self.time_out = time_out
        self.time_in: Optional[int] = None
        self.routes: List[str] = []
        self.stops = 0
        self.bags = 0
        self.km: Optional[float] = 0.0
        self.expected: Optional[float] = None
        # (time out, routes, stops, bags) for each departure in the trip.
        self.legs: List[Tuple[Optional[int], List[str], int, int]] = []

    @property
    def minutes(self) -> Optional[int]:
        if self.time_out is None or self.time_in is None:
            return None
        return self.time_in - self.time_out

    @property
    def label(self) -> str:
        if len(self.legs) < 2:
            return " + ".join(self.routes)
        first, *added = self.legs
        return " + ".join(first[1]) + "".join(
            f", adding {' + '.join(routes)} at {clock(t_out)}" for t_out, routes, _, _ in added
        )

    @property
    def biggest_load(self) -> int:
        """The most bags on the truck at once (each leg's bags leave separately)."""
        return max((bags for _, _, _, bags in self.legs), default=self.bags)

    def absorb(self, other: "Trip") -> None:
        """Add a route taken out while this trip was under way."""
        self.legs += other.legs
        self.routes += other.routes
        self.stops += other.stops
        self.bags += other.bags
        self.km = self.km + other.km if self.km is not None and other.km is not None else None
        self.time_in = max(self.time_in, other.time_in)


def _norm(header: str) -> str:
    return re.sub(r"\(s\)|[^a-z]", "", header.lower())


def minutes_of(value) -> Optional[int]:
    """Minutes after midnight for a time cell (time, datetime, day fraction or text)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        value = value.time()
    if isinstance(value, datetime.time):
        minutes = value.hour * 60 + value.minute
    elif isinstance(value, (int, float)):
        if not 0 <= value < 1:
            return None
        minutes = round(value * 24 * 60)
    else:
        m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*([ap])?\.?m?\.?\s*", str(value), re.I)
        if not m:
            return None
        hour, minute = int(m.group(1)), int(m.group(2))
        if m.group(3):
            hour = hour % 12 + (12 if m.group(3).lower() == "p" else 0)
        minutes = hour * 60 + minute
    if minutes < EARLIEST_HOUR * 60:
        minutes += 12 * 60
    return minutes


def clock(minutes: Optional[int]) -> str:
    return "" if minutes is None else f"{minutes // 60}:{minutes % 60:02d}"


def split_drivers(text: str) -> Tuple[str, ...]:
    names = re.split(r"\s*(?:/|&|,|\+|\band\b)\s*", text.strip())
    return tuple(n for n in names if n)


def load_driver_log(path: str) -> List[Dict]:
    """Rows of ``OrderDataByDeliveryRoute``: route, drivers, time_out, time_in, row."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if ROUTE_STATS_SHEET not in wb.sheetnames:
        raise SystemExit(f"{path}: no {ROUTE_STATS_SHEET} sheet.")
    columns: Optional[Dict[str, int]] = None
    log: List[Dict] = []
    for r, row in enumerate(wb[ROUTE_STATS_SHEET].iter_rows(values_only=True), start=1):
        if columns is None:
            names = {_norm(str(v)): i for i, v in enumerate(row) if v is not None}
            if {"deliveryroute", "timein", "timeout"} <= set(names) and ("driver" in names or "drivers" in names):
                columns = {
                    "route": names["deliveryroute"],
                    "driver": names.get("driver", names.get("drivers")),
                    "in": names["timein"],
                    "out": names["timeout"],
                }
            continue
        cell = lambda key: row[columns[key]] if columns[key] < len(row) else None  # noqa: E731
        route = str(cell("route") or "").strip()
        if not route or route.startswith("Grand Total"):
            continue
        log.append({
            "row": r,
            "route": route,
            "drivers": split_drivers(str(cell("driver") or "")),
            "time_out": minutes_of(cell("out")),
            "time_in": minutes_of(cell("in")),
            "raw_out": cell("out"),
            "raw_in": cell("in"),
        })
    wb.close()
    if columns is None:
        raise SystemExit(f"{path}: {ROUTE_STATS_SHEET} has no Delivery Route / Driver(s) / Time In / Time Out header.")
    return log


def route_sizes(path: str) -> Dict[str, Tuple[int, int]]:
    """{route: (stops, bags)} from the orders on ``Customers``."""
    orders, _ = load_orders(path)
    sizes: Dict[str, Tuple[int, int]] = {}
    for o in orders:
        route = o.get("Delivery Route", "")
        stops, bags = sizes.get(route, (0, 0))
        sizes[route] = (stops + 1, bags + int(float(o["Number of Bags"])))
    return sizes


def route_km(stop_order: Optional[str]) -> Dict[str, float]:
    """{route: km from the start to its last stop} from ``stop_order.json``."""
    if not stop_order or not os.path.exists(stop_order):
        return {}
    with open(stop_order, encoding="utf-8") as f:
        routes = json.load(f)["routes"]
    km: Dict[str, float] = {}
    for route, stops in routes.items():
        legs = [s.get("km_from_previous") for s in stops]
        if legs and all(leg is not None for leg in legs):
            km[route] = sum(legs)
    return km


def _orders(n: int) -> str:
    return f"{n} order" + ("" if n == 1 else "s")


def build_trips(
    log: List[Dict], sizes: Dict[str, Tuple[int, int]], km: Dict[str, float]
) -> Tuple[List[Trip], List[str]]:
    """Group the logged routes into trips; return (trips, problems in the log)."""
    problems: List[str] = []
    logged = []
    for entry in log:
        route, drivers = entry["route"], entry["drivers"]
        stops, _ = sizes.get(route, (0, 0))
        where = f"{ROUTE_STATS_SHEET} row {entry['row']}, {route}"
        if not drivers:
            if stops:
                problems.append(f"{where}: {_orders(stops)} but no driver.")
            continue
        if not stops:
            problems.append(f"{where}: driver {', '.join(drivers)} but no orders on this route.")
        for key in ("out", "in"):
            raw = entry[f"raw_{key}"]
            if raw not in (None, "") and entry[f"time_{key}"] is None:
                problems.append(f"{where}: can't read Time {key.title()} {raw!r}.")
        t_out, t_in = entry["time_out"], entry["time_in"]
        if t_out is not None and t_in is not None and t_in < t_out:
            problems.append(
                f"{where}: Time In {clock(t_in)} is before Time Out {clock(t_out)}; "
                f"taken as out {clock(t_in)}, in {clock(t_out)}."
            )
            t_out, t_in = t_in, t_out
        driver_key = tuple(sorted(d.lower() for d in drivers))
        logged.append([entry, where, driver_key, t_out, t_in])

    # A Time In with no Time Out, at or before the driver's first Time Out,
    # can't be a return: it's the Time Out typed in the wrong column.
    first_out: Dict[Tuple, int] = {}
    for _, _, driver_key, t_out, _ in logged:
        if t_out is not None:
            first_out[driver_key] = min(first_out.get(driver_key, t_out), t_out)
    for item in logged:
        _, where, driver_key, t_out, t_in = item
        if t_out is None and t_in is not None and driver_key in first_out and t_in <= first_out[driver_key]:
            problems.append(
                f"{where}: only a Time In ({clock(t_in)}), before this driver's first Time Out; "
                f"taken as Time Out {clock(t_in)}."
            )
            item[3], item[4] = t_in, None

    trips: Dict[Tuple, Trip] = {}
    for entry, where, driver_key, t_out, t_in in logged:
        route = entry["route"]
        stops, bags = sizes.get(route, (0, 0))
        key = (driver_key, t_out) if t_out is not None else (driver_key, None, route)
        trip = trips.get(key)
        if trip is None:
            trip = trips[key] = Trip(entry["drivers"], t_out)
        trip.routes.append(route)
        trip.stops += stops
        trip.bags += bags
        trip.km = trip.km + km[route] if trip.km is not None and route in km else None
        if t_in is not None:
            trip.time_in = max(trip.time_in or 0, t_in)

    for route, (stops, _) in sizes.items():
        if stops and not any(e["route"] == route for e in log):
            problems.append(f"{route}: {_orders(stops)} but no row on {ROUTE_STATS_SHEET}.")
    for trip in trips.values():
        who = ", ".join(trip.drivers)
        if trip.time_out is None:
            back = f"Time In {clock(trip.time_in)} but " if trip.time_in is not None else ""
            problems.append(f"{trip.label} ({who}): {back}no Time Out; left out of the per-minute numbers.")
        elif trip.time_in is None:
            problems.append(
                f"{trip.label} ({who}): Time Out {clock(trip.time_out)} but no Time In; "
                "left out of the per-minute numbers."
            )

    for trip in trips.values():
        trip.legs = [(trip.time_out, list(trip.routes), trip.stops, trip.bags)]

    # A route taken out before the driver's trip is back was added to that
    # trip on the way: one trip, from the first Time Out to the last Time In.
    merged: List[Trip] = []
    by_drivers: Dict[Tuple, List[Trip]] = {}
    for trip in trips.values():
        if trip.minutes is None:
            merged.append(trip)
        else:
            by_drivers.setdefault(tuple(sorted(d.lower() for d in trip.drivers)), []).append(trip)
    for driver_trips in by_drivers.values():
        driver_trips.sort(key=lambda t: t.time_out)
        current = driver_trips[0]
        for nxt in driver_trips[1:]:
            if nxt.time_out < current.time_in - OVERLAP_TOLERANCE_MIN:
                current.absorb(nxt)
            else:
                merged.append(current)
                current = nxt
        merged.append(current)
    ordered = sorted(merged, key=lambda t: (t.drivers[0].lower(), t.time_out if t.time_out is not None else 9999))
    return ordered, problems


def _solve(a: List[List[float]], b: List[float]) -> Optional[List[float]]:
    """Solve a small linear system by Gaussian elimination (None if singular)."""
    n = len(b)
    m = [row[:] + [b[i]] for i, row in enumerate(a)]
    for col in range(n):
        pivot = max(range(col, n), key=lambda r: abs(m[r][col]))
        if abs(m[pivot][col]) < 1e-9:
            return None
        m[col], m[pivot] = m[pivot], m[col]
        for r in range(n):
            if r != col:
                f = m[r][col] / m[col][col]
                m[r] = [x - f * y for x, y in zip(m[r], m[col])]
    return [m[i][n] / m[i][i] for i in range(n)]


def fit_trip_minutes(trips: Sequence[Trip]) -> Optional[Dict[str, float]]:
    """Least-squares ``minutes = base + per_stop*stops + per_bag*bags`` over the timed trips.

    A term that comes out negative (too few trips to tell) is dropped and the
    rest refitted.  Returns the coefficients and ``r2``, or None.
    """
    timed = [t for t in trips if t.minutes is not None]
    terms = ["per_stop", "per_bag"]
    while True:
        if len(timed) < len(terms) + 2:
            return None
        xs = [[1.0] + [float(t.stops if term == "per_stop" else t.bags) for term in terms] for t in timed]
        ys = [float(t.minutes) for t in timed]
        k = len(terms) + 1
        ata = [[sum(x[i] * x[j] for x in xs) for j in range(k)] for i in range(k)]
        aty = [sum(x[i] * y for x, y in zip(xs, ys)) for i in range(k)]
        coef = _solve(ata, aty)
        if coef is None:
            return None
        negative = [term for term, c in zip(terms, coef[1:]) if c < 0]
        if not negative:
            break
        terms.remove(negative[0])
    model = {"base": coef[0], "per_stop": 0.0, "per_bag": 0.0, "trips": len(timed)}
    model.update(zip(terms, coef[1:]))
    mean = sum(ys) / len(ys)
    ss_tot = sum((y - mean) ** 2 for y in ys)
    ss_res = sum((y - sum(c * x for c, x in zip(coef, row))) ** 2 for row, y in zip(xs, ys))
    model["r2"] = 1 - ss_res / ss_tot if ss_tot else 0.0
    return model


def expected_minutes(model: Dict[str, float], stops: int, bags: int) -> float:
    return model["base"] + model["per_stop"] * stops + model["per_bag"] * bags


def _union_minutes(intervals: List[Tuple[int, int]]) -> int:
    total, end = 0, None
    for start, stop in sorted(intervals):
        if end is None or start > end:
            total += stop - start
            end = stop
        elif stop > end:
            total += stop - end
            end = stop
    return total


def driver_stats(trips: Sequence[Trip]) -> List[Dict]:
    """One row per driver (see the module docstring)."""
    names: Dict[str, str] = {}
    mine: Dict[str, List[Trip]] = {}
    for trip in trips:
        for d in trip.drivers:
            names.setdefault(d.lower(), d)
            mine.setdefault(d.lower(), []).append(trip)
    rows = []
    for key, driver_trips in mine.items():
        timed = [t for t in driver_trips if t.minutes is not None]
        minutes = _union_minutes([(t.time_out, t.time_in) for t in timed])
        timed_bags = sum(t.bags for t in timed)
        timed_stops = sum(t.stops for t in timed)
        expected = sum(t.expected for t in timed if t.expected is not None)
        actual = sum(t.minutes for t in timed if t.expected is not None)
        ordered = sorted(timed, key=lambda t: t.time_out)
        gaps = [b.time_out - a.time_in for a, b in zip(ordered, ordered[1:]) if b.time_out >= a.time_in]
        rows.append({
            "driver": names[key],
            "routes": [r for t in driver_trips for r in t.routes],
            "trips": len(driver_trips),
            "timed_trips": len(timed),
            "stops": sum(t.stops for t in driver_trips),
            "bags": sum(t.bags for t in driver_trips),
            "minutes": minutes,
            "bags_per_min": timed_bags / minutes if minutes else None,
            "stops_per_min": timed_stops / minutes if minutes else None,
            "min_per_stop": minutes / timed_stops if timed_stops else None,
            "pace": expected / actual if actual else None,
            "max_bags_trip": max(t.biggest_load for t in driver_trips),
            "max_stops_trip": max(t.stops for t in driver_trips),
            "first_out": min((t.time_out for t in timed), default=None),
            "last_in": max((t.time_in for t in timed), default=None),
            "reload_min": sum(gaps) / len(gaps) if gaps else None,
        })
    return rows


def _leaders(rows: List[Dict], key: str) -> List[Dict]:
    scored = [r for r in rows if r[key] is not None]
    if not scored:
        return []
    best = max(r[key] for r in scored)
    return [r for r in scored if abs(r[key] - best) < 1e-9]


def _table(headers: Sequence[str], rows: Sequence[Sequence[str]], right: Sequence[bool]) -> None:
    widths = [max(len(h), *(len(r[i]) for r in rows)) if rows else len(h) for i, h in enumerate(headers)]
    fmt = lambda cells: "  ".join(  # noqa: E731
        c.rjust(w) if r else c.ljust(w) for c, w, r in zip(cells, widths, right)
    ).rstrip()
    print(fmt(headers))
    print(fmt(["-" * w for w in widths]))
    for row in rows:
        print(fmt(row))


def _num(value: Optional[float], digits: int = 2) -> str:
    return "" if value is None else f"{value:.{digits}f}"


def report(trips: List[Trip], drivers: List[Dict], model: Optional[Dict], problems: List[str]) -> None:
    if problems:
        print(f"DRIVER LOG PROBLEMS ({len(problems)}) - fix them in the workbook and run again:")
        for p in problems:
            print(f"  - {p}")
        print()

    print("ANSWERS")
    questions = [
        ("Most bags delivered", "bags", lambda r: f"{r['bags']} bags"),
        ("Most bags per minute", "bags_per_min", lambda r: f"{r['bags_per_min']:.2f} bags/min ({r['minutes']} min out)"),
        ("Most stops", "stops", lambda r: f"{r['stops']} stops"),
        ("Most stops per minute", "stops_per_min",
         lambda r: f"{r['stops_per_min']:.3f} stops/min = {r['min_per_stop']:.1f} min per stop ({r['minutes']} min out)"),
        ("Fastest for the size of their trips", "pace",
         lambda r: f"{r['pace'] * 100:.0f}% of the typical time for the same stops and bags"),
    ]
    for n, (question, key, describe) in enumerate(questions, start=1):
        leaders = _leaders(drivers, key)
        answer = "; ".join(f"{r['driver']}: {describe(r)}" for r in leaders) or "not enough data"
        print(f"  {n}. {question}: {answer}")
    print()

    print("DRIVERS (by bags; per-minute numbers count only trips with both times)")
    rows = sorted(drivers, key=lambda r: (-r["bags"], r["driver"].lower()))
    _table(
        ["Driver", "Trips", "Stops", "Bags", "Min out", "Bags/min", "Stops/min", "Min/stop", "Pace",
         "Biggest load", "Out", "Back", "Reload min"],
        [[
            r["driver"],
            f"{r['trips']}" + ("" if r["timed_trips"] == r["trips"] else f" ({r['timed_trips']} timed)"),
            str(r["stops"]), str(r["bags"]), str(r["minutes"] or ""),
            _num(r["bags_per_min"]), _num(r["stops_per_min"], 3), _num(r["min_per_stop"], 1),
            "" if r["pace"] is None else f"{r['pace'] * 100:.0f}%",
            f"{r['max_bags_trip']} bags", clock(r["first_out"]), clock(r["last_in"]), _num(r["reload_min"], 0),
        ] for r in rows],
        [False, True, True, True, True, True, True, True, True, True, True, True, True],
    )
    print()

    print("TRIPS (miles: start to last stop of each route, from the stop order; added up for several routes)")
    _table(
        ["Driver", "Routes", "Out", "Back", "Min", "Expected", "Stops", "Bags", "Miles"],
        [[
            ", ".join(t.drivers), t.label, clock(t.time_out), clock(t.time_in), str(t.minutes or ""),
            _num(t.expected, 0), str(t.stops), str(t.bags),
            "" if t.km is None else f"{t.km / KM_PER_MILE:.1f}",
        ] for t in trips],
        [False, False, True, True, True, True, True, True, True],
    )
    print()

    print("FOR ASSIGNING ROUTES NEXT SEASON")
    if model:
        print(
            f"  - A typical trip took {model['base']:.0f} min, plus {model['per_stop']:.1f} min per stop and "
            f"{model['per_bag'] * 10:.1f} min per 10 bags (fitted on {model['trips']} timed trips, "
            f"R² {model['r2']:.2f}). Use it to size routes: e.g. 6 stops and 100 bags ≈ "
            f"{expected_minutes(model, 6, 100):.0f} min."
        )
        print(
            "  - Pace compares each driver's time with that typical time for the same stops and bags "
            "(over 100% = faster), so it is fairer than raw bags/min between big-order and small-order routes."
        )
    print("  - Biggest load: the most bags a driver took out at once, a guide to their truck or trailer size.")
    print("  - A route taken out before the driver's trip was back was added to that trip on the way "
          "(\"adding ... at\" in TRIPS): one trip, from the first Time Out to the last Time In.")
    print("  - Out/Back: when each driver started and finished, i.e. how long they are available.")
    print("  - Reload min: average time at the lot between trips.")
    print("  - A driver with only one timed trip has a noisy rate; weigh their numbers less.")


def write_csv(directory: str, trips: List[Trip], drivers: List[Dict]) -> None:
    os.makedirs(directory, exist_ok=True)
    with open(os.path.join(directory, "drivers.csv"), "w", newline="", encoding="utf-8") as f:
        fields = ["driver", "routes", "trips", "timed_trips", "stops", "bags", "minutes", "bags_per_min",
                  "stops_per_min", "min_per_stop", "pace", "max_bags_trip", "max_stops_trip", "first_out",
                  "last_in", "reload_min"]
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in drivers:
            row = dict(r)
            row["routes"] = "; ".join(r["routes"])
            row["first_out"], row["last_in"] = clock(r["first_out"]), clock(r["last_in"])
            w.writerow({k: "" if row[k] is None else row[k] for k in fields})
    with open(os.path.join(directory, "trips.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["drivers", "routes", "time_out", "time_in", "minutes", "expected_minutes", "stops", "bags", "miles"])
        for t in trips:
            w.writerow([
                "; ".join(t.drivers), "; ".join(t.routes), clock(t.time_out), clock(t.time_in),
                "" if t.minutes is None else t.minutes, _num(t.expected, 1), t.stops, t.bags,
                "" if t.km is None else f"{t.km / KM_PER_MILE:.2f}",
            ])


def check_driver_log(path: str) -> List[str]:
    """Every route with bags needs a Time Out and a later Time In; return the problems."""
    sizes = route_sizes(path)
    log = load_driver_log(path)
    problems: List[str] = []
    for entry in log:
        stops, bags = sizes.get(entry["route"], (0, 0))
        if not bags:
            continue
        where = f"{ROUTE_STATS_SHEET} row {entry['row']}, {entry['route']} ({bags} bags)"
        t_out, t_in = entry["time_out"], entry["time_in"]
        missing = [name for name, t, raw in (("Time Out", t_out, entry["raw_out"]), ("Time In", t_in, entry["raw_in"]))
                   if t is None]
        for name in missing:
            raw = entry["raw_out" if name == "Time Out" else "raw_in"]
            problems.append(f"{where}: " + (f"can't read {name} {raw!r}." if raw not in (None, "") else f"no {name}."))
        if not missing and t_in <= t_out:
            problems.append(f"{where}: Time In {clock(t_in)} is not later than Time Out {clock(t_out)}.")
    for route, (stops, bags) in sizes.items():
        if bags and not any(e["route"] == route for e in log):
            problems.append(f"{route} ({bags} bags): no row on {ROUTE_STATS_SHEET} for its driver and times.")
    return problems


def analyze(path: str, stop_order: Optional[str] = None):
    """Return (trips, driver rows, model, problems) for a season workbook."""
    sizes = route_sizes(path)
    trips, problems = build_trips(load_driver_log(path), sizes, route_km(stop_order))
    model = fit_trip_minutes(trips)
    if model:
        for t in trips:
            if t.minutes is not None:
                t.expected = expected_minutes(model, t.stops, t.bags)
    return trips, driver_stats(trips), model, problems


def main() -> None:
    parser = argparse.ArgumentParser(description="Analyze the delivery day from the driver log.")
    parser.add_argument("--input", required=True, help="Season workbook (.xlsx) with the driver log filled in.")
    parser.add_argument("--stop-order", default="output/stop_order.json",
                        help="Stop order from route_order.py, for the miles of each trip (optional).")
    parser.add_argument("--csv-dir", help="Also write drivers.csv and trips.csv to this directory.")
    parser.add_argument("--skip-checks", action="store_true",
                        help="Report even if routes are missing times (not recommended).")
    args = parser.parse_args()
    if not args.input.lower().endswith((".xlsx", ".xlsm")):
        sys.exit("The driver log is on the OrderDataByDeliveryRoute sheet: give the .xlsx workbook.")
    failed = check_driver_log(args.input)
    if failed:
        print(f"DRIVER LOG CHECK FAILED ({len(failed)}): every route with bags needs a Time Out and a later Time In.")
        for p in failed:
            print(f"  - {p}")
        if not args.skip_checks:
            print("Stopping: fix the times in the workbook and run again (or use --skip-checks).")
            sys.exit(1)
        print("Continuing anyway (--skip-checks); the numbers guess around these.\n")
    else:
        print("Driver log check passed: every route with bags has a Time Out and a later Time In.\n")
    trips, drivers, model, problems = analyze(args.input, args.stop_order)
    report(trips, drivers, model, problems)
    if args.csv_dir:
        write_csv(args.csv_dir, trips, drivers)
        print(f"\nWrote {os.path.join(args.csv_dir, 'drivers.csv')} and trips.csv.")


if __name__ == "__main__":
    main()
