"""The season workflows end to end, on an anonymized copy of a real season.

``fixtures/season_anonymized.xlsx`` is made from the Fall 2026 workbook by
``fixtures/anonymize_workbook.py``: 182 orders on 42 routes, 3132 bags,
24 map links without a pin, 2 red comments, the summary rows, the
``BasicOrderStats`` and ``OrderDataByDeliveryRoute`` sheets and the
driver log (with the real log's mistakes), but made-up customers,
addresses, locations and drivers.

The workflows (see CLAUDE.md):

1. Check the workbook before generating (checks 1-5, small-route suggestions).
2. Order each route's stops (offline here: straight-line distances).
3. Generate the order forms and the maps, in stop order.
4. Check the generated PDFs (one page per order, each page for the right order).
5. The wrapper script ties them together and stops when a check fails.
6. Analyze the driver log after the delivery day.

Map pages are printed by Chrome from Google Maps, which the tests can't do;
they check map pages made up with the same text instead.

Run from the project root:

    python3 -m unittest discover tests
"""

import contextlib
import csv
import datetime
import io
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "src"))

import analyze_deliveries  # noqa: E402
import check_workbook  # noqa: E402
import generate_maps_pdf as maps  # noqa: E402
import generate_order_forms as forms  # noqa: E402
import route_order  # noqa: E402

SEASON = os.path.join(HERE, "fixtures", "season_anonymized.xlsx")
ORDERS, BAGS, ROUTES = 182, 3132, 42

# Customers columns (1-based, for openpyxl).
COL_STREET_ADDRESS, COL_MAP_LINK, COL_ROUTE, COL_CONDITIONAL_ROUTE, COL_BAGS = 8, 11, 12, 13, 17


def quiet(fn, *args, **kwargs):
    """Call ``fn`` and return ``(result, captured stdout)``."""
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        result = fn(*args, **kwargs)
    return result, out.getvalue()


def edited_season(tmp_dir: str, edit) -> str:
    """A copy of the season workbook changed by ``edit(wb)``."""
    import openpyxl

    path = os.path.join(tmp_dir, "edited.xlsx")
    wb = openpyxl.load_workbook(SEASON)
    edit(wb)
    wb.save(path)
    return path


def first_order_row(ws, route: str) -> int:
    return next(r for r in range(3, ws.max_row + 1)
                if ws.cell(r, COL_ROUTE).value == route and isinstance(ws.cell(r, COL_BAGS).value, int)
                and ws.cell(r, COL_BAGS).value > 0)


def straight_distances(points):
    """Metres between points as the crow flies (instead of the Routes API)."""
    def metres(a, b):
        dy = (a[0] - b[0]) * 111_000
        dx = (a[1] - b[1]) * 111_000 * math.cos(math.radians(a[0]))
        return math.hypot(dx, dy)
    return [[metres(a, b) for b in points] for a in points]


def plan_season(path: str, out_path: str):
    """Order the stops like route_order.py does, offline; return (routes, warnings)."""
    orders, _ = check_workbook.load_orders(path)
    routes, warnings = route_order.plan_routes(orders, {}, straight_distances)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"start": {"name": route_order.START_NAME, "lat": route_order.START[0],
                             "lng": route_order.START[1]}, "routes": routes}, f)
    return routes, warnings


def text_pdf(lines) -> bytes:
    """A one-page PDF with these lines of text (like a printed map page's header)."""
    def escape(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = "BT /F1 12 Tf 40 760 Td " + " ".join(f"({escape(line)}) Tj 0 -16 Td" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def write_map_pdfs(path: str, out_dir: str, stops=None, edit_pages=None) -> None:
    """Map PDFs as generate_maps_pdf.py lays them out: per route, in stop order,
    each page headed by the link's address, the route and "Stop n of N"."""
    groups = maps.group_by_label(maps.read_records(path))
    for label, group in groups.items():
        group = route_order.in_stop_order(group, lambda rec: rec[5], stops)
        pages = []
        for url, _, bags, note, _, row in group:
            stop = route_order.stop_label(row, stops)
            lines = [maps.extract_address(url) or "", label or "", f"{bags} bags"]
            lines += [f"Stop {stop}"] if stop else []
            lines += [note] if note else []
            pages.append(text_pdf(lines))
        if edit_pages:
            pages = edit_pages(label, pages)
        maps.merge_pdf_pages(pages, os.path.join(out_dir, check_workbook.route_file(label or "") + ".pdf"))


class BeforeGeneratingTests(unittest.TestCase):
    """Workflow 1: the checks just before generating."""

    def check(self, path):
        problems, _ = quiet(check_workbook.check_before, path)
        return problems

    def test_season_passes(self):
        problems, out = quiet(check_workbook.check_before, SEASON)
        self.assertEqual(problems, [])
        self.assertIn(f"Orders: {ORDERS}  Bags: {BAGS}  With Map Link: {ORDERS}", out)

    def test_orders_and_route_stats(self):
        orders, totals = check_workbook.load_orders(SEASON)
        self.assertEqual((len(orders), totals), (ORDERS, str(BAGS)))
        self.assertEqual(len({o["Delivery Route"] for o in orders}), ROUTES)
        stats = check_workbook.load_route_stats(SEASON)
        self.assertEqual(stats["Fairfax 01A"], ("14", "227"))
        self.assertEqual(stats["Fairfax 03B"], ("0", "0"))

    def test_small_route_suggestions(self):
        orders, _ = check_workbook.load_orders(SEASON)
        suggestions = check_workbook.route_suggestions(orders)
        small = [s.split(" has ")[0] for s in suggestions]
        self.assertEqual(small, ["Fairfax 06B", "Fairfax 08C", "Fairfax 12B", "Fairfax Station 04",
                                 "Fairfax Station 09", "George Mason 04", "George Mason 05B", "George Mason 06"])
        self.assertTrue(all("Outlier" not in s for s in suggestions))

    def test_missing_conditional_route_formula(self):
        # The Fall 2026 problem: two rows without the column M formula made
        # OrderDataByDeliveryRoute show 13 orders on Fairfax 01A instead of 14.
        def edit(wb):
            ws = wb["Customers"]
            r = first_order_row(ws, "Fairfax 01A")
            ws.cell(r, COL_CONDITIONAL_ROUTE).value = None
            stats = wb["OrderDataByDeliveryRoute"]
            row = next(r for r in range(1, stats.max_row + 1) if stats.cell(r, 2).value == "Fairfax 01A")
            stats.cell(row, 3).value = 13
        with tempfile.TemporaryDirectory() as tmp:
            problems = self.check(edited_season(tmp, edit))
        self.assertEqual(len(problems), 2)
        self.assertIn("Conditional Delivery route", problems[0])
        self.assertIn("copy its formula from the row above", problems[0])
        self.assertIn("Fairfax 01A", problems[1])
        self.assertIn("13", problems[1])

    def test_bags_changed_after_the_totals(self):
        def edit(wb):
            ws = wb["Customers"]
            r = first_order_row(ws, "Burke 03")
            ws.cell(r, COL_BAGS).value += 5
        with tempfile.TemporaryDirectory() as tmp:
            problems = self.check(edited_season(tmp, edit))
        self.assertTrue(any("Totals" in p for p in problems), problems)
        self.assertTrue(any("BasicOrderStats" in p for p in problems), problems)
        self.assertTrue(any("Burke 03" in p for p in problems), problems)

    def test_map_link_problems(self):
        def edit(wb):
            ws = wb["Customers"]
            ws.cell(first_order_row(ws, "Burke 01"), COL_MAP_LINK).value = None
            other = ws.cell(first_order_row(ws, "Burke 04"), COL_MAP_LINK).value
            ws.cell(first_order_row(ws, "Burke 03"), COL_MAP_LINK).value = other
        with tempfile.TemporaryDirectory() as tmp:
            problems = self.check(edited_season(tmp, edit))
        self.assertTrue(any("(Burke 01): no Map Link." in p for p in problems), problems)
        self.assertTrue(any("(Burke 03): Map Link is for" in p for p in problems), problems)


class StopOrderTests(unittest.TestCase):
    """Workflow 2: each route's stops ordered from the Sideburn Run lot."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.stop_file = os.path.join(cls.tmp.name, "stop_order.json")
        cls.routes, cls.warnings = plan_season(SEASON, cls.stop_file)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_every_order_is_one_stop_on_its_route(self):
        orders, _ = check_workbook.load_orders(SEASON)
        self.assertEqual(check_workbook._check_stop_order(self.stop_file, orders), [])
        self.assertEqual(len(self.routes), ROUTES)
        self.assertEqual(sum(len(s) for s in self.routes.values()), ORDERS)

    def test_links_without_a_pin_use_the_map_centre(self):
        self.assertEqual(len(self.warnings), 24)
        self.assertTrue(all("no pin found; used the map's centre instead." in w for w in self.warnings))

    def test_each_route_is_the_shortest_pass(self):
        # Small routes are solved exactly: no other order is shorter.
        import itertools

        for route, stops in self.routes.items():
            if len(stops) > 6:
                continue
            points = [route_order.START] + [(s["lat"], s["lng"]) for s in stops]
            dist = straight_distances(points)
            length = route_order.path_length(dist, list(range(len(stops))))
            best = min(route_order.path_length(dist, list(p)) for p in itertools.permutations(range(len(stops))))
            self.assertAlmostEqual(length, best, places=6, msg=route)

    def test_stop_labels(self):
        stops = route_order.load_stop_order(self.stop_file)
        first = self.routes["Fairfax 01A"][0]["row"]
        self.assertEqual(route_order.stop_label(first, stops), "1 of 14")

    def test_no_api_key(self):
        result = subprocess.run(
            [sys.executable, os.path.join(ROOT, "src", "route_order.py"), "--input", SEASON,
             "--output", os.path.join(self.tmp.name, "unused.json")],
            capture_output=True, text=True, env={k: v for k, v in os.environ.items() if k != "GOOGLE_MAPS_API_KEY"},
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("No Google Maps API key", result.stdout)


class GenerateAndCheckTests(unittest.TestCase):
    """Workflows 3-4: order forms and maps for the whole season, then the checks."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.stop_file = os.path.join(cls.tmp.name, "stop_order.json")
        plan_season(SEASON, cls.stop_file)
        cls.stops = route_order.load_stop_order(cls.stop_file)
        cls.records, cls.form_output = quiet(forms.parse_order_records, SEASON)
        cls.orders_dir = os.path.join(cls.tmp.name, "orders")
        quiet(forms.save_order_forms, cls.records, cls.orders_dir, cls.stops)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_order_form_records(self):
        self.assertEqual(len(self.records), ORDERS)
        self.assertEqual(sum(int(r["bags"]) for r in self.records), BAGS)
        church = [r for r in self.records if r["customer"].startswith("Hillcrest Community Church")]
        self.assertEqual(len(church), 1)
        self.assertEqual(church[0]["amount_support"], "")
        self.assertEqual(sum(1 for r in self.records if not r["order_no"]), 27)  # mail-in orders etc.

    def test_one_order_form_pdf_per_route(self):
        from PyPDF2 import PdfReader

        pdfs = sorted(os.listdir(self.orders_dir))
        self.assertEqual(len(pdfs), ROUTES)
        pages = {name: len(PdfReader(os.path.join(self.orders_dir, name)).pages) for name in pdfs}
        self.assertEqual(sum(pages.values()), ORDERS)
        self.assertEqual(pages["Fairfax_01A.pdf"], 14)

    def test_order_forms_pass_the_checks(self):
        problems, _ = quiet(check_workbook.check_after, SEASON, None, None, self.orders_dir, self.stop_file)
        self.assertEqual(problems, [])

    def test_checks_catch_order_forms_in_the_wrong_order(self):
        # The PDF-merging bug: pages mixed up while the page counts still matched.
        orders_dir = os.path.join(self.tmp.name, "orders_mixed_up")
        shutil.copytree(self.orders_dir, orders_dir)
        quiet(forms.save_order_forms, self.records, orders_dir)  # spreadsheet order, no stop numbers
        problems, _ = quiet(check_workbook.check_after, SEASON, None, None, orders_dir, self.stop_file)
        self.assertTrue(problems)
        self.assertTrue(all(p.startswith("order forms: ") and "but isn't" in p for p in problems), problems)

    def test_map_records(self):
        records = maps.read_records(SEASON)
        self.assertEqual(len(records), ORDERS)
        notes = sorted(rec[3] for rec in records if rec[3])
        self.assertEqual(len(notes), 2)  # exactly the red comments
        self.assertEqual(len(maps.group_by_label(records)), ROUTES)
        for url, _, _, _, street, _ in records:
            self.assertTrue(check_workbook.same_street_address(maps.extract_address(url), street), url)

    def test_map_pages_pass_the_checks(self):
        maps_dir = os.path.join(self.tmp.name, "maps")
        os.makedirs(maps_dir)
        write_map_pdfs(SEASON, maps_dir, self.stops)
        problems, _ = quiet(check_workbook.check_after, SEASON, None, maps_dir, None, self.stop_file)
        self.assertEqual(problems, [])

    def test_checks_catch_map_pages_in_the_wrong_order(self):
        maps_dir = os.path.join(self.tmp.name, "maps_swapped")
        os.makedirs(maps_dir)

        def swap_first_two(label, pages):
            return [pages[1], pages[0]] + pages[2:] if label == "Fairfax 01A" else pages
        write_map_pdfs(SEASON, maps_dir, self.stops, swap_first_two)
        problems, _ = quiet(check_workbook.check_after, SEASON, None, maps_dir, None, self.stop_file)
        # Each page is for the wrong order and has the wrong stop number.
        self.assertEqual({p.split(" should be ")[0].split(" doesn't say")[0] for p in problems},
                         {"maps: Fairfax_01A.pdf page 1", "maps: Fairfax_01A.pdf page 2"})
        self.assertEqual(len(problems), 4)

    def test_checks_catch_a_map_page_printed_twice(self):
        maps_dir = os.path.join(self.tmp.name, "maps_twice")
        os.makedirs(maps_dir)

        def repeat_first(label, pages):
            return [pages[0], pages[0]] + pages[2:] if label == "Burke 03" else pages
        write_map_pdfs(SEASON, maps_dir, self.stops, repeat_first)
        problems, _ = quiet(check_workbook.check_after, SEASON, None, maps_dir, None, self.stop_file)
        self.assertEqual(problems, [
            problems[0],
            "maps: Burke_03.pdf page 2 doesn't say 'Stop 2 of 4'.",
        ])
        self.assertTrue(problems[0].startswith("maps: Burke_03.pdf page 2 should be"), problems)

    def test_checks_catch_a_missing_stop_number(self):
        maps_dir = os.path.join(self.tmp.name, "maps_unnumbered")
        os.makedirs(maps_dir)
        write_map_pdfs(SEASON, maps_dir, None)  # printed without the stop order
        problems, _ = quiet(check_workbook.check_after, SEASON, None, maps_dir, None, self.stop_file)
        self.assertTrue(any("doesn't say 'Stop" in p for p in problems), problems)


class WrapperScriptTests(unittest.TestCase):
    """Workflow 5: ./run_generate_maps_pdf.sh (order forms only: maps need Chrome)."""

    WRAPPER = os.path.join(ROOT, "run_generate_maps_pdf.sh")

    def run_wrapper(self, *args, key=None):
        env = {k: v for k, v in os.environ.items() if k != "GOOGLE_MAPS_API_KEY"}
        if key:
            env["GOOGLE_MAPS_API_KEY"] = key
        return subprocess.run(["bash", self.WRAPPER, *args], capture_output=True, text=True, env=env,
                              stdin=subprocess.DEVNULL, timeout=600)

    def test_order_forms_without_route_order(self):
        key = "test-key-must-not-be-shown"
        with tempfile.TemporaryDirectory() as out_dir:
            result = self.run_wrapper("--orders", "--no-route-order", SEASON, out_dir, key=key)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(len([f for f in os.listdir(out_dir) if f.endswith(".pdf")]), ROUTES)
            self.assertFalse(os.path.exists(os.path.join(out_dir, "stop_order.json")))
        self.assertIn("Not ordering the stops", result.stdout)
        self.assertIn("After generating: all checks passed.", result.stdout)
        self.assertNotIn(key, result.stdout + result.stderr)

    def test_stops_when_a_check_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            def edit(wb):
                ws = wb["Customers"]
                ws.cell(first_order_row(ws, "Fairfax 01A"), COL_CONDITIONAL_ROUTE).value = None
            path = edited_season(tmp, edit)
            out_dir = os.path.join(tmp, "out")
            result = self.run_wrapper("--orders", "--no-route-order", path, out_dir)
            self.assertEqual(result.returncode, 5)
            self.assertFalse(os.path.exists(out_dir) and any(f.endswith(".pdf") for f in os.listdir(out_dir)))
        self.assertIn("Stopping: fix the workbook", result.stderr)

    def test_usage_errors(self):
        self.assertEqual(self.run_wrapper("--bogus", SEASON).returncode, 2)
        self.assertEqual(self.run_wrapper("--orders").returncode, 2)
        self.assertEqual(self.run_wrapper("--orders", "no/such/workbook.xlsx").returncode, 3)


class AfterDeliveryTests(unittest.TestCase):
    """Workflow 6: the driver log analysis (driver names are anonymized)."""

    @classmethod
    def setUpClass(cls):
        (cls.trips, drivers, cls.model, cls.problems), _ = quiet(analyze_deliveries.analyze, SEASON)
        cls.drivers = {d["driver"]: d for d in drivers}

    # The driver log as first entered in Fall 2026, before it was fixed:
    # {cell: value} on OrderDataByDeliveryRoute (H = Time In, I = Time Out).
    FIRST_LOG = {"I7": None, "H9": datetime.time(8, 25), "I9": datetime.time(9, 5),
                 "H31": datetime.time(7, 30), "I31": None, "H63": datetime.time(8, 20), "I63": None}
    FIRST_LOG_FAILURES = [
        "OrderDataByDeliveryRoute row 7, Burke 02B (75 bags): no Time Out.",
        "OrderDataByDeliveryRoute row 9, Burke 04 (63 bags): Time In 8:25 is not later than Time Out 9:05.",
        "OrderDataByDeliveryRoute row 31, Fairfax 12B (12 bags): no Time Out.",
        "OrderDataByDeliveryRoute row 63, Outlier (77 bags): no Time Out.",
    ]

    def first_log(self, tmp, extra=None):
        def edit(wb):
            for cell, value in {**self.FIRST_LOG, **(extra or {})}.items():
                wb["OrderDataByDeliveryRoute"][cell].value = value
        return edited_season(tmp, edit)

    def run_script(self, path, *args):
        return subprocess.run(
            [sys.executable, os.path.join(ROOT, "src", "analyze_deliveries.py"), "--input", path,
             "--stop-order", "none", *args], capture_output=True, text=True)

    def test_driver_log_check_passes(self):
        # Routes with no bags (e.g. Fairfax 03B) need no times.
        self.assertEqual(analyze_deliveries.check_driver_log(SEASON), [])

    def test_driver_log_check_catches_the_first_log(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(analyze_deliveries.check_driver_log(self.first_log(tmp)), self.FIRST_LOG_FAILURES)

    def test_driver_log_check_catches_equal_times_and_missing_time_in(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self.first_log(tmp, {"H5": datetime.time(9, 20), "H40": None})
            failures = analyze_deliveries.check_driver_log(path)
        self.assertIn("OrderDataByDeliveryRoute row 5, Burke 01 (41 bags): "
                      "Time In 9:20 is not later than Time Out 9:20.", failures)
        self.assertIn("OrderDataByDeliveryRoute row 40, Fairfax Station 02 (128 bags): no Time In.", failures)

    def test_script_stops_when_the_check_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self.first_log(tmp)
            stopped = self.run_script(path)
            skipped = self.run_script(path, "--skip-checks")
        self.assertEqual(stopped.returncode, 1)
        self.assertIn("DRIVER LOG CHECK FAILED (4)", stopped.stdout)
        self.assertNotIn("ANSWERS", stopped.stdout)
        self.assertEqual(skipped.returncode, 0, skipped.stderr)
        self.assertIn("Continuing anyway (--skip-checks)", skipped.stdout)
        self.assertIn("ANSWERS", skipped.stdout)

    def test_script_reports_when_the_check_passes(self):
        result = self.run_script(SEASON)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Driver log check passed", result.stdout)
        self.assertIn("1. Most bags delivered:", result.stdout)

    def test_no_log_problems(self):
        self.assertEqual(self.problems, [])

    def test_routes_added_on_the_way(self):
        # Three drivers added a route to a trip that was under way.
        added = [t.label for t in self.trips if ", adding " in t.label]
        self.assertEqual(sorted(added), ["Fairfax 01A, adding Fairfax 02B at 9:47",
                                         "Fairfax 11, adding George Mason 01 at 9:15",
                                         "Fairfax 12A + Fairfax 12B, adding Fairfax 02A at 9:30"])
        trip = next(t for t in self.trips if t.label.startswith("Fairfax 01A"))
        self.assertEqual((trip.minutes, trip.stops, trip.bags, trip.biggest_load), (200, 20, 311, 227))

    def test_skip_checks_guesses_around_the_first_log(self):
        with tempfile.TemporaryDirectory() as tmp:
            (_, _, _, problems), _ = quiet(analyze_deliveries.analyze, self.first_log(tmp))
        self.assertEqual(len(problems), 5)
        self.assertIn("Burke 04: Time In 8:25 is before Time Out 9:05", problems[0])
        self.assertIn("Fairfax 12B: only a Time In (7:30)", problems[1])
        self.assertEqual(sum("left out of the per-minute numbers" in p for p in problems), 2)

    def test_totals(self):
        self.assertEqual(sum(t.bags for t in self.trips), BAGS)
        self.assertEqual(sum(t.stops for t in self.trips), ORDERS)
        self.assertEqual(len(self.drivers), 9)

    def test_answers(self):
        def leader(key):
            return [r["driver"] for r in analyze_deliveries._leaders(list(self.drivers.values()), key)]
        top = leader("bags")
        self.assertEqual(self.drivers[top[0]]["bags"], 473)
        self.assertEqual(leader("bags_per_min"), top)
        self.assertEqual(leader("stops"), top)
        self.assertEqual(leader("stops_per_min"), top)
        self.assertEqual(self.drivers[top[0]]["bags_per_min"], 1.6)  # 473 / 295 = 1.603...
        self.assertEqual(self.drivers[top[0]]["stops_per_min"], 0.1)  # 30 / 295 = 0.1016...
        self.assertEqual(self.drivers[top[0]]["minutes"], 295)

    def test_trip_model(self):
        self.assertEqual(self.model["trips"], 26)
        self.assertGreater(self.model["per_stop"], 0)
        self.assertGreater(self.model["per_bag"], 0)

    def test_report_and_csv(self):
        drivers = list(self.drivers.values())
        _, out = quiet(analyze_deliveries.report, self.trips, drivers, self.model, self.problems)
        self.assertNotIn("DRIVER LOG PROBLEMS", out)
        self.assertIn("1. Most bags delivered:", out)
        with tempfile.TemporaryDirectory() as tmp:
            analyze_deliveries.write_csv(tmp, self.trips, drivers)
            with open(os.path.join(tmp, "drivers.csv"), newline="") as f:
                rows = list(csv.DictReader(f))
            with open(os.path.join(tmp, "trips.csv"), newline="") as f:
                trips = list(csv.DictReader(f))
        self.assertEqual(len(rows), 9)
        self.assertEqual(sum(int(t["bags"]) for t in trips), BAGS)


if __name__ == "__main__":
    unittest.main()
