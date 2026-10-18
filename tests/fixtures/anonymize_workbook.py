#!/usr/bin/env python3
"""Make an anonymized copy of a season workbook for the tests.

Usage (from the project root)::

    .venv/bin/python tests/fixtures/anonymize_workbook.py "input/<season workbook>.xlsx" \
        tests/fixtures/season_anonymized.xlsx

The copy keeps what the scripts and checks depend on: the sheet layout, the
routes, neighbourhoods, bag counts, money, the summary rows, the
``BasicOrderStats`` and ``OrderDataByDeliveryRoute`` values and the driver
log times, which comments are red, and which map links have a pin.

Everything that could identify a customer or a driver is replaced:
names, emails, house numbers and streets (and the map links, rebuilt for
the made-up address), comments, delivery instructions, order numbers and
driver names.  Each map location is moved to a random point near its
route's area (rounded to ~1 km), at least 100 m from every real customer
location.  Made-up names and streets that match a word in any name or
street on the sheet are not used.  Only
the order rows, a sample of past customers and the summary rows are kept.
Values are written instead of formulas (the scripts read the values).

Before saving, it checks that no real name, email, street, comment,
instruction, order number or driver name appears anywhere in the copy,
and stops if one does.  The script is deterministic (fixed seed).
"""

import math
import random
import re
import sys
import warnings
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font

HERE = sys.path[0]
sys.path.insert(0, f"{HERE}/../../src")

from generate_maps_pdf import extract_coordinates  # noqa: E402
from route_order import pin_coordinates  # noqa: E402
from sheet_reader import _is_red  # noqa: E402

LAST_COL = 23  # Customers columns A..W
D, E, F, G, H, I, K, L, O, Q, S, V, W = 4, 5, 6, 7, 8, 9, 11, 12, 15, 17, 19, 22, 23
PAST_CUSTOMERS_KEPT = 30
MIN_METRES_FROM_REAL = 100  # every made-up location is at least this far from every real one
SUMMARY_LABELS = ("Totals", "Pallets", "Total Order", "Extra Bags", "Total Received")

LAST_NAMES = [
    "Ashgrove", "Brackenridge", "Cobbleworth", "Dunmorrow", "Elderwick", "Fenwhistle", "Gildersleeve",
    "Hollowmere", "Ivybrook", "Juniperlund", "Kestrelton", "Larkspurr", "Mossgrave", "Nettlefold",
    "Oakhaven", "Pinecrestor", "Quillstone", "Rookwood", "Sablethorn", "Thistlewood", "Umberfield",
    "Valemont", "Wexborough", "Yarrowby", "Zephyrine", "Amberlake", "Bramblecote", "Cindervale",
    "Driftwell", "Emberton", "Foxglover", "Glimmerby", "Heathcombe", "Inglenook", "Jasperwick",
]
FIRST_NAMES = [
    "Avery", "Dana", "Eden", "Finley", "Harper", "Indigo", "Jordan",
    "Kendall", "Logan", "Morgan", "Noel", "Oakley", "Parker", "Quinn", "Reese", "Sage", "Taylor",
    "Rowan", "Skyler", "Emery", "Hayden", "Jessie", "Kai", "Lane", "Marlowe", "Peyton", "Remy",
]
STREET_NAMES = [
    "Acorn Glade", "Bluebell", "Copperleaf", "Dewberry", "Evergleam", "Fiddlehead", "Goldfinch",
    "Hazelnook", "Ironwood Glen", "Jaywing", "Kingfisher Bend", "Lilypad", "Meadowlark Rise",
    "Nightjar", "Orchard Mist", "Pebblebrook", "Quailrun", "Redstart", "Snowberry", "Tanager",
    "Upland Fern", "Vireo", "Willowherb", "Yellowwood Knoll", "Zinnia", "Alderbloom", "Bobolink",
    "Cloudberry", "Dragonfly", "Elfin Oak",
]
SUFFIXES = [("Drive", "Dr"), ("Court", "Ct"), ("Road", "Rd"), ("Lane", "Ln"), ("Place", "Pl"),
            ("Street", "St"), ("Way", "Way"), ("Terrace", "Ter")]
INSTRUCTIONS = [
    "Left side of the driveway.", "In front of the garage door.", "By the mailbox, please.",
    "Stack along the front walk.", "Behind the side gate, on the patio.", "Right side of the driveway, near the fence.",
    "10 bags in the front bed, the rest by the garage.", "At the top of the driveway, please don't block the garage.",
    "Next to the shed in the back yard.", "On the grass strip by the curb.",
]
RED_NOTES = ["Need one pallet", "Call before delivery", "Deliver after 10am"]
OTHER_COMMENTS = ["Paid by check", "Same as last year", "Cash at pickup"]


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def location(link: str) -> Optional[Tuple[float, float]]:
    return pin_coordinates(link) or extract_coordinates(link)


def metres(a: Tuple[float, float], b: Tuple[float, float]) -> float:
    dy = (a[0] - b[0]) * 111_000
    dx = (a[1] - b[1]) * 111_000 * math.cos(math.radians(a[0]))
    return math.hypot(dx, dy)


def map_link(number: str, street: str, abbreviation: str, town: str, zip_code, where: Tuple[float, float],
             pin: bool) -> str:
    place = f"{number}+{street.replace(' ', '+')}+{abbreviation},+{town.replace(' ', '+')},+VA+{zip_code}"
    lat, lng = where
    link = f"https://www.google.com/maps/place/{place}/@{lat:.7f},{lng:.7f},17z/data="
    if pin:
        return link + f"!4m6!3m5!1s0x0:0x0!8m2!3d{lat:.7f}!4d{lng:.7f}!16s"
    return link + "!4m2!3m1!1s0x0:0x0"


def is_summary(ws, r: int) -> bool:
    return any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip().startswith(SUMMARY_LABELS)
               for c in range(1, LAST_COL + 1))


def anonymize(src: str, dst: str) -> None:
    rng = random.Random(1865)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.worksheets[0]

    header_row = next(r for r in range(1, 11) if ws.cell(r, D).value == "LastName")
    orders, past, summary = [], [], []
    for r in range(header_row + 1, ws.max_row + 1):
        bags = ws.cell(r, Q).value
        if isinstance(bags, (int, float)) and bags > 0 and (ws.cell(r, D).value or ws.cell(r, H).value):
            orders.append(r)
        elif is_summary(ws, r):
            summary.append(r)
        elif ws.cell(r, D).value:
            past.append(r)
    past = past[:: max(1, len(past) // PAST_CUSTOMERS_KEPT)][:PAST_CUSTOMERS_KEPT]

    # Made-up names and streets that are no one's name or street on the sheet.
    real_words = {w for r in range(header_row + 1, ws.max_row + 1) for c in (D, E, G)
                  for w in re.findall(r"[a-z]+", str(ws.cell(r, c).value or "").lower())}
    last_names = [n for n in LAST_NAMES if n.lower() not in real_words]
    first_names = [n for n in FIRST_NAMES if n.lower() not in real_words]
    streets = [n for n in STREET_NAMES if not set(n.lower().split()) & real_words]
    real_places = [p for p in (location(str(ws.cell(r, K).value or "")) for r in range(header_row + 1, ws.max_row + 1)) if p]

    # Each route's area, rounded to ~1 km.
    points: Dict[str, List[Tuple[float, float]]] = {}
    for r in orders:
        where = location(str(ws.cell(r, K).value or ""))
        if where:
            points.setdefault(str(ws.cell(r, L).value), []).append(where)
    area = {route: (round(sum(p[0] for p in pts) / len(pts), 2), round(sum(p[1] for p in pts) / len(pts), 2))
            for route, pts in points.items()}
    default_area = (38.81, -77.31)

    out = openpyxl.Workbook()
    new = out.active
    new.title = ws.title
    for r in range(1, header_row + 1):
        for c in range(1, LAST_COL + 1):
            new.cell(r, c, ws.cell(r, c).value)

    used_addresses = set()
    order_no = 50000
    sensitive: List[str] = []
    out_row = header_row
    for r in orders + past + [None] + summary:
        out_row += 1
        if r is None:
            continue  # a blank row before the summary rows, like the real sheet
        for c in range(1, LAST_COL + 1):
            src_cell, cell = ws.cell(r, c), new.cell(out_row, c)
            cell.value = src_cell.value
            cell.number_format = src_cell.number_format
        if r in summary:
            continue
        for c in (D, E, F, G, H, K, O, V, W):
            if ws.cell(r, c).value not in (None, ""):
                sensitive.append(str(ws.cell(r, c).value))

        last, first = rng.choice(last_names), rng.choice(first_names)
        if r in orders and not ws.cell(r, S).value:
            last, first = "Hillcrest Community Church", ""  # the benefactor: no support amount
        # (cell(r, c, None) wouldn't clear the copied value)
        new.cell(out_row, D).value = last
        new.cell(out_row, E).value = first or None
        while True:
            number = str(rng.randint(3000, 12999))
            street, (suffix, abbreviation) = rng.choice(streets), rng.choice(SUFFIXES)
            if (number, street) not in used_addresses:
                used_addresses.add((number, street))
                break
        new.cell(out_row, F, int(number) if isinstance(ws.cell(r, F).value, int) else number)
        new.cell(out_row, G, f"{street} {suffix}")
        new.cell(out_row, H, f"{number} {street} {suffix}")
        old_link = str(ws.cell(r, K).value or "")
        if old_link:
            centre = area.get(str(ws.cell(r, L).value), default_area)
            for _ in range(1000):
                where = (centre[0] + rng.uniform(-0.004, 0.004), centre[1] + rng.uniform(-0.004, 0.004))
                if all(metres(where, p) >= MIN_METRES_FROM_REAL for p in real_places):
                    break
            else:
                sys.exit(f"No made-up location {MIN_METRES_FROM_REAL} m from every customer near {centre}.")
            new.cell(out_row, K, map_link(number, street, abbreviation, str(ws.cell(r, I).value or "Fairfax"),
                                          ws.cell(r, 10).value or "", where, pin=pin_coordinates(old_link) is not None))
        if ws.cell(r, O).value:
            new.cell(out_row, O, f"{first or 'office'}.{last.split()[0]}{out_row}@example.com".lower())

        comment, comment_cell = ws.cell(r, V).value, new.cell(out_row, V)
        if comment not in (None, ""):
            text = str(comment).strip()
            if re.fullmatch(r"Order \d+", text):
                order_no += 1
                comment_cell.value = f"Order {order_no}"
            elif _is_red(ws.cell(r, V).font.color if ws.cell(r, V).font else None):
                comment_cell.value = rng.choice(RED_NOTES)
                comment_cell.font = Font(color="FFFF0000")
            else:
                comment_cell.value = rng.choice(OTHER_COMMENTS)
        instructions = ws.cell(r, W).value
        if instructions not in (None, "") and str(instructions).strip().upper() != "TBD":
            new.cell(out_row, W, rng.choice(INSTRUCTIONS))

    driver_col = None
    for name in ("BasicOrderStats", "OrderDataByDeliveryRoute"):
        if name not in wb.sheetnames:
            continue
        copy = out.create_sheet(name)
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.value is not None:
                    copy.cell(cell.row, cell.column, cell.value).number_format = cell.number_format
    if "OrderDataByDeliveryRoute" in out.sheetnames:
        log = out["OrderDataByDeliveryRoute"]
        driver_col = next((c.column for row in log.iter_rows(max_row=10) for c in row
                           if isinstance(c.value, str) and c.value.strip().lower().startswith("driver")), None)
        aliases: Dict[str, str] = {}
        if driver_col:
            for r in range(1, log.max_row + 1):
                cell = log.cell(r, driver_col)
                if isinstance(cell.value, str) and not cell.value.strip().lower().startswith("driver"):
                    sensitive.append(cell.value)
                    names = [n for n in re.split(r"\s*(?:/|&|,|\+|\band\b)\s*", cell.value.strip()) if n]
                    for n in names:
                        aliases.setdefault(n.lower(), f"Driver {chr(ord('A') + len(aliases))}")
                    cell.value = " / ".join(aliases[n.lower()] for n in names)

    kept = [str(new.cell(r, c).value) for r in range(1, header_row + 1) for c in range(1, LAST_COL + 1)]
    kept += [str(new.cell(r, c).value) for r in range(header_row + 1, new.max_row + 1) for c in (L, L + 1, L + 2)]
    kept += [str(ws.cell(r, c).value) for r in summary for c in range(1, LAST_COL + 1)]
    kept += [str(cell.value) for sheet in out.worksheets[1:] for row in sheet.iter_rows() for cell in row
             if not (driver_col and sheet.title == "OrderDataByDeliveryRoute" and cell.column == driver_col)]
    _check_no_leaks(out, sensitive, kept)
    out.save(dst)
    print(f"Wrote {dst}: {len(orders)} orders, {len(past)} past customers, {len(summary)} summary rows.")


def _check_no_leaks(wb, sensitive: List[str], kept: List[str]) -> None:
    """Stop if any real value, or a real word of 4+ letters from one, is in the copy.

    House numbers alone aren't checked (a made-up one may equal a real one on
    another street); whole addresses are.  Words in ``kept`` (headers, routes,
    neighbourhoods, summary and stats sheets, copied on purpose) and the
    made-up texts are allowed.
    """
    phrases, words = set(), set()
    for value in sensitive:
        value = norm(value)
        if not value.isdigit():
            (phrases if re.search(r"[ @/]", value) else words).add(value)
        if "@" not in value and "http" not in value:
            words.update(re.findall(r"[a-z]{4,}", value))
    # Words the copy uses for its own reasons (street types, towns, routes, common words).
    allowed = {w for value in INSTRUCTIONS + RED_NOTES + OTHER_COMMENTS for w in re.findall(r"[a-z]{4,}", value.lower())}
    allowed |= {s.lower() for pair in SUFFIXES for s in pair}
    allowed |= {"fairfax", "station", "burke", "clifton", "springfield", "order", "please", "https", "google",
                "maps", "place", "community", "church", "hillcrest", "office", "example", "driver"}
    allowed |= {w for value in kept for w in re.findall(r"[a-z]{4,}", value.lower())}
    words -= allowed
    made_up = [norm(v) for v in INSTRUCTIONS + RED_NOTES + OTHER_COMMENTS]
    phrases = {p for p in phrases if not any(p in m for m in made_up)}
    phrases -= {norm(v) for v in kept}
    leaks = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                text = norm(cell.value)
                if words & set(re.findall(r"[a-z0-9']+", text)) or any(p in text for p in phrases):
                    leaks.append(f"{ws.title}!{cell.coordinate}")
    if leaks:
        sys.exit(f"Not saved: real values found in {len(leaks)} cells: {', '.join(leaks[:20])}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    anonymize(sys.argv[1], sys.argv[2])
