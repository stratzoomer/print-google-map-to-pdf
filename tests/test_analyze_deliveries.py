"""Tests for the post-delivery analysis (src/analyze_deliveries.py).

The workbook is the made-up season fixture (``build_season_workbook``) with a
made-up driver log added to ``OrderDataByDeliveryRoute``.  Its routes:
Fairfax 01A (2 stops, 11 bags), Fairfax 07 (2 stops, 30 bags), Fairfax 12A
(1 stop, 40 bags) and Fairfax 13 (no orders).
"""

import datetime
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", "src"))

import analyze_deliveries as ad  # noqa: E402
import route_order  # noqa: E402
from test_sheet_parsing import build_season_workbook, quiet  # noqa: E402

T = datetime.time


def with_driver_log(path: str, log) -> None:
    """Add Driver(s)/Time In/Time Out columns: log = {route: (driver, in, out)}."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb["OrderDataByDeliveryRoute"]
    ws.cell(2, 7, "Driver(s):")
    ws.cell(2, 8, "Time In")
    ws.cell(2, 9, "Time Out")
    for r in range(4, ws.max_row + 1):
        entry = log.get(ws.cell(r, 2).value)
        if entry:
            for c, value in enumerate(entry, start=7):
                ws.cell(r, c, value)
    wb.save(path)


class AnalyzeDeliveriesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = os.path.join(self.tmp.name, "season.xlsx")
        build_season_workbook(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def analyze(self, log):
        with_driver_log(self.xlsx, log)
        (trips, drivers, model, problems), _ = quiet(ad.analyze, self.xlsx)
        return trips, {d["driver"]: d for d in drivers}, problems

    def test_driver_totals_and_rates(self):
        trips, drivers, problems = self.analyze({
            "Fairfax 01A": ("Ann", T(9, 0), T(8, 0)),
            "Fairfax 07": ("Ann", T(9, 40), T(9, 10)),
            "Fairfax 12A": ("Bo", T(10, 20), T(10, 0)),
        })
        self.assertEqual(problems, [])
        ann, bo = drivers["Ann"], drivers["Bo"]
        self.assertEqual((ann["trips"], ann["stops"], ann["bags"], ann["minutes"]), (2, 4, 41, 90))
        self.assertEqual(ann["bags_per_min"], 0.45)  # 0.4555...
        self.assertEqual(ann["stops_per_hour"], 2.66)  # 4 stops in 90 min = 2.666...
        self.assertEqual(ann["min_per_stop"], 22.5)
        self.assertEqual(ann["reload_min"], 10)
        self.assertEqual((bo["bags"], bo["minutes"], bo["max_bags_trip"]), (40, 20, 40))

    def test_routes_out_together_are_one_trip(self):
        trips, drivers, _ = self.analyze({
            "Fairfax 01A": ("Ann", T(9, 0), T(8, 0)),
            "Fairfax 07": ("Ann", T(9, 30), T(8, 0)),
        })
        self.assertEqual([(t.label, t.minutes, t.stops, t.bags) for t in trips], [("Fairfax 01A + Fairfax 07", 90, 4, 41)])
        self.assertEqual(drivers["Ann"]["max_bags_trip"], 41)

    def test_log_problems(self):
        _, drivers, problems = self.analyze({
            "Fairfax 01A": ("Ann", T(8, 0), T(9, 0)),  # in before out
            "Fairfax 07": ("Bo", T(9, 0), None),  # no out
            "Fairfax 13": ("Cy", T(9, 0), T(8, 0)),  # no orders
        })
        self.assertIn("Time In 8:00 is before Time Out 9:00; taken as out 8:00, in 9:00.", problems[0])
        self.assertEqual(drivers["Ann"]["minutes"], 60)
        self.assertTrue(any("Fairfax 07 (Bo): Time In 9:00 but no Time Out" in p for p in problems))
        self.assertIsNone(drivers["Bo"]["bags_per_min"])
        self.assertTrue(any("Fairfax 13: driver Cy but no orders" in p for p in problems))
        self.assertTrue(any("Fairfax 12A: 1 order but no driver." in p for p in problems))

    def test_time_in_before_first_time_out_is_a_time_out(self):
        trips, _, problems = self.analyze({
            "Fairfax 01A": ("Ann", T(9, 0), T(8, 0)),
            "Fairfax 07": ("Ann", T(8, 0), None),
        })
        self.assertTrue(any("taken as Time Out 8:00" in p for p in problems))
        self.assertEqual([t.label for t in trips], ["Fairfax 01A + Fairfax 07"])

    def test_route_added_on_the_way_joins_the_trip(self):
        trips, drivers, problems = self.analyze({
            "Fairfax 01A": ("Ann", T(10, 0), T(8, 0)),
            "Fairfax 07": ("Ann", T(10, 0), T(9, 0)),  # out before Fairfax 01A is back
            "Fairfax 12A": ("Ann", T(11, 0), T(10, 30)),  # a new trip
        })
        self.assertEqual(problems, [])
        self.assertEqual(
            [(t.label, t.time_out, t.time_in, t.stops, t.bags) for t in trips],
            [("Fairfax 01A, adding Fairfax 07 at 9:00", 480, 600, 4, 41), ("Fairfax 12A", 630, 660, 1, 40)],
        )
        ann = drivers["Ann"]
        self.assertEqual((ann["trips"], ann["minutes"]), (2, 150))
        self.assertEqual(ann["max_bags_trip"], 40)  # 30 and 11 bags left separately

    def test_rounded_times_are_not_an_added_route(self):
        trips, _, _ = self.analyze({
            "Fairfax 01A": ("Ann", T(9, 25), T(8, 0)),
            "Fairfax 07": ("Ann", T(10, 0), T(9, 20)),  # 5 minutes before: a new trip
        })
        self.assertEqual([t.label for t in trips], ["Fairfax 01A", "Fairfax 07"])

    def test_shared_trip_counts_for_each_driver(self):
        _, drivers, _ = self.analyze({"Fairfax 07": ("Ann / Bo", T(9, 0), T(8, 0))})
        self.assertEqual((drivers["Ann"]["bags"], drivers["Bo"]["bags"]), (30, 30))


class DriverLogCheckTests(unittest.TestCase):
    """Every route with bags needs a Time Out and a later Time In."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = os.path.join(self.tmp.name, "season.xlsx")
        build_season_workbook(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def check(self, log):
        with_driver_log(self.xlsx, log)
        return quiet(ad.check_driver_log, self.xlsx)[0]

    def test_complete_log_passes(self):
        # Fairfax 13 has no bags, so it needs no times.
        self.assertEqual(self.check({
            "Fairfax 01A": ("Ann", T(9, 0), T(8, 0)),
            "Fairfax 07": ("Ann", "9:40", "9:10"),
            "Fairfax 12A": ("Bo", T(10, 20), T(10, 0)),
        }), [])

    def test_missing_flipped_equal_and_unreadable_times(self):
        self.assertEqual(self.check({
            "Fairfax 01A": ("Ann", T(8, 0), T(9, 0)),
            "Fairfax 07": ("Ann", T(9, 40), T(9, 40)),
            "Fairfax 12A": ("Bo", "soon", None),
        }), [
            "OrderDataByDeliveryRoute row 4, Fairfax 01A (11 bags): Time In 8:00 is not later than Time Out 9:00.",
            "OrderDataByDeliveryRoute row 5, Fairfax 07 (30 bags): Time In 9:40 is not later than Time Out 9:40.",
            "OrderDataByDeliveryRoute row 6, Fairfax 12A (40 bags): no Time Out.",
            "OrderDataByDeliveryRoute row 6, Fairfax 12A (40 bags): can't read Time In 'soon'.",
        ])

    def test_route_with_bags_missing_from_the_sheet(self):
        import openpyxl

        wb = openpyxl.load_workbook(self.xlsx)
        wb["OrderDataByDeliveryRoute"].delete_rows(6)  # Fairfax 12A
        wb.save(self.xlsx)
        failures = self.check({
            "Fairfax 01A": ("Ann", T(9, 0), T(8, 0)),
            "Fairfax 07": ("Ann", T(9, 40), T(9, 10)),
        })
        self.assertEqual(failures, ["Fairfax 12A (40 bags): no row on OrderDataByDeliveryRoute for its driver and times."])


class HelperTests(unittest.TestCase):
    def test_minutes_of(self):
        self.assertEqual(ad.minutes_of(T(9, 5)), 545)
        self.assertEqual(ad.minutes_of(datetime.datetime(2026, 10, 4, 13, 20)), 800)
        self.assertEqual(ad.minutes_of(0.5), 720)
        self.assertEqual(ad.minutes_of("10:45"), 645)
        self.assertEqual(ad.minutes_of("1:15 pm"), 795)
        self.assertEqual(ad.minutes_of("1:15"), 795)  # no am/pm before 6:00: afternoon
        self.assertIsNone(ad.minutes_of("soon"))
        self.assertIsNone(ad.minutes_of(None))

    def test_round_down(self):
        self.assertEqual(ad.round_down(1.6033), 1.6)
        self.assertEqual(ad.round_down(0.0999), 0.09)
        self.assertEqual(ad.round_down(1.15), 1.15)  # not 1.14 (1.15 * 100 = 114.999...)
        self.assertEqual(ad.round_down(26.199), 26.19)

    def test_split_drivers(self):
        self.assertEqual(ad.split_drivers("Ann / Bo & Cy and Di"), ("Ann", "Bo", "Cy", "Di"))
        self.assertEqual(ad.split_drivers(""), ())

    def test_fit_with_miles(self):
        trips = []
        for stops, bags, miles in [(2, 20, 3), (5, 40, 9), (4, 100, 4), (8, 90, 12), (3, 60, 30), (6, 150, 7), (5, 50, 15)]:
            t = ad.Trip(("Ann",), 0)
            t.stops, t.bags, t.miles = stops, bags, miles
            t.time_in = 10 + 3 * stops + 0.5 * bags + 2 * miles
            trips.append(t)
        model = ad.fit_trip_minutes(trips)
        self.assertTrue(model["uses_miles"])
        for term, value in (("base", 10), ("per_stop", 3), ("per_bag", 0.5), ("per_mile", 2)):
            self.assertAlmostEqual(model[term], value, places=6, msg=term)
        self.assertAlmostEqual(ad.expected_minutes(model, 6, 100, 5), 10 + 18 + 50 + 10)
        # A trip without miles: they are left out for all.
        trips[0].miles = None
        self.assertFalse(ad.fit_trip_minutes(trips)["uses_miles"])

    def test_trip_miles(self):
        start, a, b = route_order.START, (38.82, -77.31), (38.83, -77.31)
        key = route_order._cache_key
        trip = ad.Trip(("Ann",), 480)
        trip.legs = [(480, ["Route A"], 1, 10), (540, ["Route B"], 1, 10)]  # Route B added on the way
        stops = {"Route A": [a], "Route B": [b]}
        known = {key(start, a): {"m": 1609.344, "s": 60}, key(b, start): {"m": 3218.688, "s": 120}}
        miles, legs, estimated = ad.trip_miles(trip, stops, known)
        # Lot -> A (1 mile), A -> B (estimated: straight line x DETOUR), B -> lot (2 miles).
        a_to_b = ad._straight_metres(a, b) * ad.DETOUR / ad.METRES_PER_MILE
        self.assertAlmostEqual(miles, 3 + a_to_b)
        self.assertEqual((legs, estimated), (3, 1))
        self.assertEqual(ad.trip_miles(trip, {"Route A": [a]}, known), (None, 0, 0))

    def test_fit_recovers_exact_model(self):
        trips = []
        for stops, bags in [(2, 20), (5, 40), (4, 100), (8, 90), (3, 60), (6, 150)]:
            t = ad.Trip(("Ann",), 0)
            t.stops, t.bags, t.time_in = stops, bags, round(10 + 3 * stops + 0.5 * bags)
            trips.append(t)
        model = ad.fit_trip_minutes(trips)
        self.assertAlmostEqual(model["base"], 10, places=6)
        self.assertAlmostEqual(model["per_stop"], 3, places=6)
        self.assertAlmostEqual(model["per_bag"], 0.5, places=6)
        self.assertAlmostEqual(model["r2"], 1)


if __name__ == "__main__":
    unittest.main()
