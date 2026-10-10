"""Tests for reading the season spreadsheet in both generator scripts.

``fixtures/season_export.csv`` is a small made-up copy of the Google Sheets
export format: a "Fall 2026" title row above the header, cp1252 encoding
(smart quote), past customers with no bags, a zero-bag row, routes that are
not adjacent, a mail-in order with a special-action comment, the benefactor
order with "TBD" instructions, an order with no map link, a blank row and
the Totals/Pallets summary rows at the bottom.  The ``.xlsx`` tests build
a workbook from the same rows (see ``build_season_workbook``).

Run from the project root:

    python3 -m unittest discover tests
"""

import contextlib
import csv
import io
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "src"))

import addresses  # noqa: E402
import check_workbook  # noqa: E402
import generate_maps_pdf as maps  # noqa: E402
import generate_order_forms as forms  # noqa: E402
import route_order  # noqa: E402

SEASON_CSV = os.path.join(HERE, "fixtures", "season_export.csv")

import sheet_reader  # noqa: E402

# Column indexes in the season sheet.
BAGS, SUPPORT, COMMENT, ZIP = 16, 18, 21, 9


def quiet(fn, *args, **kwargs):
    """Call ``fn`` and return ``(result, captured stdout)``."""
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        result = fn(*args, **kwargs)
    return result, out.getvalue()


def write_temp_csv(text: str) -> str:
    f = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, newline="", encoding="utf-8")
    with f:
        f.write(text)
    return f.name


class OrderFormParsingTests(unittest.TestCase):
    def setUp(self):
        self.records, self.output = quiet(forms.parse_order_records, SEASON_CSV)

    def test_only_rows_with_bags_are_orders(self):
        # Skips the title row, past customers with no bags, the zero-bag
        # row, the blank row and the Totals/Pallets summary rows.
        self.assertEqual(
            [r["customer"] for r in self.records],
            [
                "Testerson, Alice",
                "Mailer, Bob",
                "Testerson, Erin",
                "Church, Benefactor",
                "Nolink, Frank",
            ],
        )

    def test_fields_come_from_header_below_title_row(self):
        alice = self.records[0]
        self.assertEqual(alice["order_no"], "14812")
        self.assertEqual(alice["amount_support"], "$53.00")
        self.assertEqual(alice["city"], "Fairfax")
        self.assertEqual(alice["address"], "100 Maple Drive")
        self.assertEqual(alice["email"], "alice@example.com")
        self.assertEqual(alice["bags"], "20")
        self.assertEqual(alice["route"], "Fairfax 07")

    def test_cp1252_smart_quote_is_decoded(self):
        self.assertEqual(self.records[0]["instructions"], "End of driveway, don’t block gate")

    def test_ids_are_sequential_among_orders(self):
        self.assertEqual([r["id"] for r in self.records], [1, 2, 3, 4, 5])

    def test_mail_in_order_has_no_order_number(self):
        self.assertEqual(self.records[1]["order_no"], "")

    def test_benefactor_has_no_support_amount(self):
        self.assertEqual(self.records[3]["amount_support"], "")

    def test_warns_about_tbd_instructions_and_missing_map_link(self):
        self.assertIn("Church, Benefactor (Fairfax 12A): Delivery Instructions are still 'TBD'", self.output)
        self.assertIn("Nolink, Frank (Fairfax 01A): no Map Link", self.output)
        self.assertEqual(self.output.count("WARNING"), 2)

    def test_header_on_first_row_still_works(self):
        # Older exports (e.g. Spring 2026) had no title row.
        path = write_temp_csv(
            "LastName,FirstName,Street Address,Town,Map Link,Delivery Route,"
            "EmailAddress,Number of Bags,Support Troop Amount,Comment,Delivery Instructions\n"
            "Smith,Jo,1 Main St,Fairfax,https://maps.example/1,Fairfax 15A,"
            "jo@example.com,20, $53.00 ,Order 123,By the garage\n"
        )
        self.addCleanup(os.remove, path)
        records, _ = quiet(forms.parse_order_records, path)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["customer"], "Smith, Jo")
        self.assertEqual(records[0]["order_no"], "123")


class IsOrderRowTests(unittest.TestCase):
    def test_positive_bags_with_name(self):
        self.assertTrue(forms.is_order_row({"Number of Bags": " 5 ", "LastName": "Smith"}))

    def test_positive_bags_with_address_only(self):
        self.assertTrue(forms.is_order_row({"Number of Bags": "5", "Street Address": "1 Main St"}))

    def test_rejects_blank_zero_and_non_integer_bags(self):
        for bags in ("", "0", "38.4", "abc"):
            with self.subTest(bags=bags):
                self.assertFalse(forms.is_order_row({"Number of Bags": bags, "LastName": "Smith"}))

    def test_rejects_summary_row_without_name_or_address(self):
        self.assertFalse(forms.is_order_row({"Number of Bags": "1730", "EmailAddress": "Totals"}))


class OrderFormOutputTests(unittest.TestCase):
    def test_one_pdf_per_route_with_one_page_per_order(self):
        from PyPDF2 import PdfReader

        records, _ = quiet(forms.parse_order_records, SEASON_CSV)
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir)
            pages = {
                name: len(PdfReader(os.path.join(out_dir, name)).pages)
                for name in os.listdir(out_dir)
            }
        self.assertEqual(
            pages,
            {"Fairfax_07.pdf": 2, "Fairfax_01A.pdf": 2, "Fairfax_12A.pdf": 1},
        )


class MapRecordTests(unittest.TestCase):
    def setUp(self):
        self.records = maps.read_records(SEASON_CSV)

    def test_only_orders_with_map_links(self):
        # Same orders as the forms, minus the one with no Map Link.
        self.assertEqual([r[2] for r in self.records], ["20", "5", "10", "40"])
        self.assertTrue(all(r[0].startswith("https://www.google.com/maps/place/") for r in self.records))

    def test_route_labels(self):
        self.assertEqual(
            [r[1] for r in self.records],
            ["Fairfax 07", "Fairfax 01A", "Fairfax 07", "Fairfax 12A"],
        )

    def test_only_special_comments_become_notes(self):
        self.assertEqual([r[3] for r in self.records], [None, "Need one pallet", None, None])

    def test_sheet_row_numbers(self):
        self.assertEqual([r[5] for r in self.records], [3, 4, 7, 8])

    def test_street_addresses(self):
        self.assertEqual(
            [r[4] for r in self.records],
            ["100 Maple Drive", "200 Oak Court", "500 Maple Drive", "600 Ox Road"],
        )

    def test_max_records(self):
        self.assertEqual(len(maps.read_records(SEASON_CSV, max_records=2)), 2)

    def test_positional_format_without_header(self):
        path = write_temp_csv("https://maps.example/1,Fairfax 15A,20\nhttps://maps.example/2,Burke 01\n")
        self.addCleanup(os.remove, path)
        self.assertEqual(
            maps.read_records(path),
            [
                ("https://maps.example/1", "Fairfax 15A", "20", None, None, 1),
                ("https://maps.example/2", "Burke 01", None, None, None, 2),
            ],
        )


class GroupByLabelTests(unittest.TestCase):
    def test_non_adjacent_rows_for_a_route_are_grouped_together(self):
        groups = maps.group_by_label(maps.read_records(SEASON_CSV))
        self.assertEqual(list(groups), ["Fairfax 07", "Fairfax 01A", "Fairfax 12A"])
        self.assertEqual([r[2] for r in groups["Fairfax 07"]], ["20", "10"])


class SpecialNoteTests(unittest.TestCase):
    def test_order_numbers_and_blanks_are_not_notes(self):
        for comment in ("", "   ", "Order 14812", " Order  14812 "):
            with self.subTest(comment=comment):
                self.assertIsNone(maps.special_note(comment))

    def test_other_comments_are_notes(self):
        self.assertEqual(maps.special_note(" Need one pallet "), "Need one pallet")
        self.assertEqual(maps.special_note("Order 14812 - call first"), "Order 14812 - call first")


def build_season_workbook(path: str) -> None:
    """Write the fixture rows as an .xlsx like the real season workbook.

    Numbers are stored as numbers (money with a "$" format), the first sheet
    is ``Customers`` with another sheet after it, and the comments carry
    colours: Bob's "Need one pallet" is red, Erin's comment is "Order 14813"
    followed by a red "call before delivery", and the church has a comment
    in the default colour.
    """
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font

    with open(SEASON_CSV, newline="", encoding="cp1252") as f:
        rows = list(csv.reader(f))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            if value == "":
                continue
            cell = ws.cell(r, c)
            text = value.strip()
            if r > 2 and c - 1 in (BAGS, ZIP) and text.isdigit():
                cell.value = int(text)
            elif r > 2 and c - 1 == SUPPORT and text.startswith("$"):
                cell.value = float(text.strip("$").replace(",", ""))
                cell.number_format = '_("$"* #,##0.00_)'
            else:
                cell.value = value
    for r in range(3, ws.max_row + 1):
        last = ws.cell(r, 4).value
        comment = ws.cell(r, COMMENT + 1)
        if last == "Mailer":
            comment.font = Font(color="FFFF0000")
        elif last == "Testerson" and ws.cell(r, 5).value == "Erin":
            comment.value = CellRichText(
                "Order 14813 ", TextBlock(InlineFont(color="FFFF0000"), "call before delivery")
            )
        elif last == "Church":
            comment.value = "Benefactor"
    wb.create_sheet("Totals")["A1"] = "not the customer list"
    stats = wb.create_sheet("BasicOrderStats")
    stats.append(["Basic Order Stats"])
    stats.append(["Total Bags", 81])
    stats.append(["Orders", 5])
    stats.append(["Average Order", 16.2])
    wb.save(path)


class XlsxInputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = os.path.join(cls.tmp.name, "Mulch Sales - Fall 2026.xlsx")
        build_season_workbook(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_order_forms_match_csv(self):
        from_csv, _ = quiet(forms.parse_order_records, SEASON_CSV)
        from_xlsx, output = quiet(forms.parse_order_records, self.xlsx)
        # Same orders and fields; numbers render as in the CSV export.
        for rec in from_csv:
            rec.pop("order_no")  # Erin's comment differs in the workbook
        for rec in from_xlsx:
            rec.pop("order_no")
        self.assertEqual(from_xlsx, from_csv)
        self.assertEqual(from_xlsx[0]["amount_support"], "$53.00")
        self.assertEqual(from_xlsx[0]["bags"], "20")
        self.assertEqual(output.count("WARNING"), 2)

    def test_map_notes_are_only_the_red_text(self):
        records = maps.read_records(self.xlsx)
        self.assertEqual([r[2] for r in records], ["20", "5", "10", "40"])
        self.assertEqual(
            [r[3] for r in records],
            [None, "Need one pallet", "call before delivery", None],
        )

    def test_named_sheet(self):
        self.assertEqual(len(maps.read_records(self.xlsx, sheet="Customers")), 4)
        with self.assertRaises(KeyError):
            maps.read_records(self.xlsx, sheet="No Such Sheet")


class FormatValueTests(unittest.TestCase):
    def test_numbers(self):
        money = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
        self.assertEqual(sheet_reader._format_value(13.25, money), "$13.25")
        self.assertEqual(sheet_reader._format_value(10260, money), "$10,260.00")
        self.assertEqual(sheet_reader._format_value(20, "General"), "20")
        self.assertEqual(sheet_reader._format_value(20.0, "General"), "20")
        self.assertEqual(sheet_reader._format_value(38.4, "General"), "38.4")
        self.assertEqual(sheet_reader._format_value(None, "General"), "")


class OrderFormPagesTests(unittest.TestCase):
    """Every page of a route's PDF is the form for that page's order.

    Merging single-page PDFs with PyPDF2 used to repeat some orders and
    drop others once a route had several orders.
    """

    def test_many_orders_in_one_route(self):
        from PIL import ImageChops
        from PyPDF2 import PdfReader

        # 30 orders: enough that the old merging reliably got pages wrong.
        records = [
            {"order_no": str(14800 + i), "amount_support": f"${2.65 * (i + 4):.2f}",
             "customer": f"Customer{i}, Test", "city": "Fairfax", "address": f"{100 + i} Elm Street",
             "email": f"c{i}@example.com", "bags": str(i + 4), "route": "Fairfax 01A",
             "id": i + 1, "instructions": f"Instructions for order {i}"}
            for i in range(30)
        ]
        fonts = forms.load_fonts()
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir)
            pages = PdfReader(os.path.join(out_dir, "Fairfax_01A.pdf")).pages
            self.assertEqual(len(pages), 30)
            images = [check_workbook._page_image(p) for p in pages]
        for i, (rec, image) in enumerate(zip(records, images)):
            expected = forms.draw_order_form(rec, fonts).convert("L")
            diff = ImageChops.difference(expected, image.resize(expected.size))
            changed = sum(diff.point(lambda v: 255 if v > 100 else 0).histogram()[255:])
            with self.subTest(page=i + 1):
                self.assertLessEqual(changed, 50)


class SameStreetAddressTests(unittest.TestCase):
    def test_abbreviations_and_town(self):
        self.assertTrue(addresses.same_street_address("100 Maple Ct, Fairfax, VA 22030", "100 Maple Court"))
        self.assertTrue(addresses.same_street_address("200 Oak Hill Dr", "200 Oak Hill Drive"))
        self.assertTrue(addresses.same_street_address("300 O’Hara Rd", "300 O'Hara Road"))
        self.assertTrue(addresses.same_street_address("400 Elm Blvd", "400 Elm Boulevard"))

    def test_different_addresses(self):
        self.assertFalse(addresses.same_street_address("900 Birch Terrace, Fairfax", "300 O'Hara Rd"))
        self.assertFalse(addresses.same_street_address("100 Maple Ct", "105 Maple Court"))
        self.assertFalse(addresses.same_street_address("", ""))


class CheckWorkbookTests(unittest.TestCase):
    """The checks run before and after generating (see check_workbook.py)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = os.path.join(cls.tmp.name, "season.xlsx")
        build_season_workbook(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def check_before(self, path):
        problems, _ = quiet(check_workbook.check_before, path)
        return problems

    def edited_workbook(self, edit):
        import openpyxl

        path = os.path.join(self.tmp.name, "edited.xlsx")
        wb = openpyxl.load_workbook(self.xlsx)
        edit(wb)
        wb.save(path)
        return path

    def test_totals_and_stats_match(self):
        # The only problem in the fixture is the order with no Map Link.
        self.assertEqual(self.check_before(self.xlsx), ["700 Birch Way (Fairfax 01A): no Map Link."])
        self.assertEqual(self.check_before(SEASON_CSV), ["700 Birch Way (Fairfax 01A): no Map Link."])

    def test_totals_row_mismatch(self):
        def edit(wb):
            for row in wb["Customers"].iter_rows():
                if row[14].value == "Totals":
                    row[16].value = 80
        problems = self.check_before(self.edited_workbook(edit))
        self.assertIn("Totals row says 80 bags, but the orders add up to 81.", problems)

    def test_stats_mismatch(self):
        def edit(wb):
            wb["BasicOrderStats"]["B2"] = 90
            wb["BasicOrderStats"]["B3"] = 6
        problems = self.check_before(self.edited_workbook(edit))
        self.assertIn("BasicOrderStats 'Total Bags' is 90, but the orders add up to 81.", problems)
        self.assertIn("BasicOrderStats 'Orders' is 6, but there are 5 orders.", problems)

    def test_map_link_for_another_address(self):
        def edit(wb):
            ws = wb["Customers"]
            for r in range(3, ws.max_row + 1):
                if ws.cell(r, 4).value == "Mailer":
                    ws.cell(r, 11).value = "https://www.google.com/maps/place/900+Birch+Terrace,+Fairfax,+VA+22030/@38.8,-77.3,17z"
        problems = self.check_before(self.edited_workbook(edit))
        self.assertIn(
            "200 Oak Court (Fairfax 01A): Map Link is for '900 Birch Terrace, Fairfax, VA 22030'.",
            problems,
        )

    def test_after_generating(self):
        records, _ = quiet(forms.parse_order_records, self.xlsx)
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir)
            problems, _ = quiet(check_workbook.check_after, self.xlsx, None, None, out_dir)
            self.assertEqual(problems, [])
            os.remove(os.path.join(out_dir, "Fairfax_07.pdf"))
            problems, _ = quiet(check_workbook.check_after, self.xlsx, None, None, out_dir)
        self.assertIn("order forms: 3 pages, but there are 5 orders.", problems)
        self.assertIn("order forms: Fairfax_07.pdf has 0 page(s), expected 2.", problems)

    def test_after_generating_catches_pages_for_the_wrong_order(self):
        records, _ = quiet(forms.parse_order_records, self.xlsx)
        fonts = forms.load_fonts()
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir)
            # Rewrite Fairfax 07 with Alice's form twice instead of Alice, Erin.
            alice = forms.draw_order_form(records[0], fonts)
            alice.save(os.path.join(out_dir, "Fairfax_07.pdf"), save_all=True, append_images=[alice])
            problems, _ = quiet(check_workbook.check_after, self.xlsx, None, None, out_dir)
        self.assertEqual(
            problems,
            ["order forms: Fairfax_07.pdf page 2 should be Testerson, Erin, 500 Maple Drive (10 bags) but isn't."],
        )



def pin_link(street: str, lat: float, lng: float) -> str:
    """A made-up Google Maps link with a place pin."""
    return (f"https://www.google.com/maps/place/{street.replace(' ', '+')}/@{lat},{lng},17z"
            f"/data=!4m6!3m5!1s0x0:0x0!8m2!3d{lat}!4d{lng}!16s")


def flat_distances(points):
    """A distance matrix for made-up points: plain (scaled) Euclidean distance."""
    return [[1000 * ((a[0] - b[0]) ** 2 + (a[1] - b[1]) ** 2) ** 0.5 for b in points] for a in points]


class RouteOrderTests(unittest.TestCase):
    def test_pin_coordinates(self):
        self.assertEqual(route_order.pin_coordinates(pin_link("100 Maple Ct", 38.81, -77.31)), (38.81, -77.31))
        # The map's centre (@) is not the pin.
        self.assertIsNone(route_order.pin_coordinates(
            "https://www.google.com/maps/place/100+Maple+Ct/@38.81,-77.31,17z"))
        self.assertIsNone(route_order.pin_coordinates(""))

    def test_stops_along_a_street_are_driven_from_the_near_end(self):
        points = [(0, x) for x in (3, 0, 4, 1, 2)]
        self.assertEqual(route_order.best_order(flat_distances([(0, -1)] + points)), [1, 3, 4, 0, 2])
        self.assertEqual(route_order.best_order(flat_distances([(0, 5)] + points)), [2, 0, 4, 3, 1])

    def test_one_way_streets(self):
        # Start 0, stops 1 and 2.  1 -> 2 is short, 2 -> 1 is a long way round.
        dist = [[0, 5, 5], [5, 0, 1], [5, 9, 0]]
        self.assertEqual(route_order.best_order(dist), [0, 1])
        dist = [[0, 5, 5], [5, 0, 9], [5, 1, 0]]
        self.assertEqual(route_order.best_order(dist), [1, 0])

    def test_shortest_order_matches_trying_every_order(self):
        import itertools
        import random

        rnd = random.Random(1865)
        for n in range(1, 8):
            # Different each way, like real driving distances.
            dist = [[0 if a == b else rnd.uniform(1, 10) for b in range(n + 1)] for a in range(n + 1)]
            order = route_order.best_order(dist)
            self.assertEqual(sorted(order), list(range(n)))
            shortest = min(route_order.path_length(dist, p) for p in itertools.permutations(range(n)))
            self.assertAlmostEqual(route_order.path_length(dist, order), shortest)

    def test_long_routes_get_every_stop_once(self):
        import random

        rnd = random.Random(7)
        dist = flat_distances([(rnd.random(), rnd.random()) for _ in range(21)])
        order = route_order.best_order(dist)
        self.assertEqual(sorted(order), list(range(20)))
        self.assertLessEqual(
            route_order.path_length(dist, order),
            route_order.path_length(dist, route_order._nearest_neighbour(dist)),
        )

    def test_plan_routes(self):
        orders, _ = check_workbook.load_orders(SEASON_CSV)
        # The fixture's links have no pins; give 500 Maple Drive one, next to the start.
        pins = {orders[2]["Map Link"]: (38.8100, -77.3129)}
        routes, warnings = route_order.plan_routes(orders, pins, flat_distances)
        self.assertEqual({r: [s["row"] for s in stops] for r, stops in routes.items()},
                         {"Fairfax 07": [7, 3], "Fairfax 01A": [4, 9], "Fairfax 12A": [8]})
        self.assertIn("700 Birch Way (Fairfax 01A): no location; put last on the route.", warnings)
        self.assertIn("100 Maple Drive (Fairfax 07): no pin found; used the map's centre instead.", warnings)
        self.assertLess(routes["Fairfax 07"][0]["km_from_previous"], 0.01)


class DrivingMatrixTests(unittest.TestCase):
    """Driving distances from the Routes API (the API itself is faked)."""

    POINTS = [(38.80, -77.30), (38.81, -77.31), (38.82, -77.32)]

    def fake_api(self, calls):
        def request(origins, destinations, api_key):
            calls.append((len(origins), len(destinations)))
            elements = []
            for i, o in enumerate(origins):
                for j, d in enumerate(destinations):
                    el = {"condition": "ROUTE_EXISTS", "duration": "60s",
                          "distanceMeters": int(round(abs(o[0] - d[0]) * 100000 + (5 if i > j else 0)))}
                    # The API leaves out fields that are 0.
                    if i:
                        el["originIndex"] = i
                    if j:
                        el["destinationIndex"] = j
                    if el["distanceMeters"] == 0:
                        del el["distanceMeters"]
                    elements.append(el)
            return elements
        return request

    def test_matrix_and_cache(self):
        from unittest import mock

        calls, cache = [], {}
        with mock.patch.object(route_order, "_request_matrix", self.fake_api(calls)):
            dist = route_order.driving_matrix(self.POINTS, "test-key", cache)
            self.assertEqual(dist[0][2], 2000.0)
            self.assertEqual(dist[2][0], 2005.0)  # longer the other way
            self.assertEqual(dist[1][1], 0.0)
            self.assertEqual(len(calls), 1)
            # Everything is cached now: no second request.
            self.assertEqual(route_order.driving_matrix(self.POINTS, "test-key", cache), dist)
            self.assertEqual(len(calls), 1)

    def test_large_routes_are_split_into_requests(self):
        from unittest import mock

        calls = []
        points = [(38.80 + i / 1000, -77.30) for i in range(30)]
        with mock.patch.object(route_order, "_request_matrix", self.fake_api(calls)):
            dist = route_order.driving_matrix(points, "test-key", {})
        self.assertEqual(sorted(calls), [(5, 5), (5, 25), (25, 5), (25, 25)])
        self.assertEqual(dist[29][0], 2905.0)

    def test_no_route(self):
        from unittest import mock

        with mock.patch.object(route_order, "_request_matrix", return_value=[]):
            with self.assertRaises(route_order.RouteOrderError):
                route_order.driving_matrix(self.POINTS, "test-key", {})

    def test_refused_key_is_not_shown(self):
        import urllib.error
        from unittest import mock

        body = io.BytesIO(b'[{"error": {"code": 400, "message": "API key not valid: secret-key-123"}}]')
        error = urllib.error.HTTPError(route_order.ROUTES_API_URL, 400, "Bad Request", {}, body)
        with mock.patch("urllib.request.urlopen", side_effect=error):
            with self.assertRaises(route_order.RouteOrderError) as raised:
                route_order._request_matrix(self.POINTS[:1], self.POINTS[1:], "secret-key-123")
        self.assertIn("HTTP 400): API key not valid", str(raised.exception))
        self.assertNotIn("secret-key-123", str(raised.exception))

    def test_request(self):
        import json
        from unittest import mock

        response = mock.MagicMock()
        response.__enter__.return_value.read.return_value = b'[{"condition": "ROUTE_EXISTS"}]'
        with mock.patch("urllib.request.urlopen", return_value=response) as urlopen:
            self.assertEqual(route_order._request_matrix(self.POINTS[:1], self.POINTS[1:], "k"),
                             [{"condition": "ROUTE_EXISTS"}])
        request = urlopen.call_args[0][0]
        self.assertEqual(request.get_header("X-goog-api-key"), "k")
        body = json.loads(request.data)
        self.assertEqual(body["travelMode"], "DRIVE")
        self.assertEqual(len(body["origins"]), 1)
        self.assertEqual(len(body["destinations"]), 2)
        self.assertEqual(body["origins"][0]["waypoint"]["location"]["latLng"], {"latitude": 38.80, "longitude": -77.30})


class RouteSuggestionTests(unittest.TestCase):
    """Small routes that might be folded into an adjacent route."""

    @staticmethod
    def order(route, street, bags, lat, lng):
        return {"Delivery Route": route, "Street Address": street, "Number of Bags": str(bags),
                "Map Link": pin_link(street, lat, lng)}

    def test_small_route_lists_its_family_and_closest_route(self):
        orders = [
            self.order("Fairfax 01A", "100 Maple Ct", 5, 38.810, -77.310),
            self.order("Fairfax 01B", "200 Oak Ct", 20, 38.811, -77.310),
            self.order("Fairfax 01B", "210 Oak Ct", 20, 38.812, -77.310),
            self.order("Fairfax 01B", "220 Oak Ct", 20, 38.813, -77.310),
            self.order("Fairfax 02", "300 Elm St", 30, 38.900, -77.310),
            self.order("Fairfax 02", "310 Elm St", 30, 38.901, -77.310),
            self.order("Fairfax 02", "320 Elm St", 30, 38.902, -77.310),
        ]
        self.assertEqual(check_workbook.route_suggestions(orders), [
            "Fairfax 01A has 1 order(s), 5 bags (100 Maple Ct). "
            "Same area: Fairfax 01B (3 order(s), 60 bags). "
            "Closest route: Fairfax 01B (3 order(s), 60 bags), 0.1 km away."
        ])

    def test_big_orders_and_the_outlier_route_are_not_suggested(self):
        orders = [
            self.order("Fairfax 15A", "500 Birch Rd", 50, 38.80, -77.30),  # one big order
            self.order("Outlier", "600 Far Rd", 5, 38.90, -77.40),
            self.order("Burke 01", "700 Pine Ct", 10, 38.80, -77.29),
            self.order("Burke 01", "710 Pine Ct", 10, 38.80, -77.29),
            self.order("Burke 01", "720 Pine Ct", 10, 38.80, -77.29),
        ]
        self.assertEqual(check_workbook.route_suggestions(orders), [])

    def test_never_suggests_the_outlier_route(self):
        orders = [
            self.order("Burke 02", "800 Ash Ct", 5, 38.800, -77.300),
            self.order("Outlier", "810 Ash Ct", 5, 38.800, -77.300),
            self.order("Burke 03", "900 Fir Ct", 10, 38.810, -77.300),
            self.order("Burke 03", "910 Fir Ct", 10, 38.810, -77.300),
            self.order("Burke 03", "920 Fir Ct", 10, 38.810, -77.300),
        ]
        [suggestion] = check_workbook.route_suggestions(orders)
        self.assertIn("Closest route: Burke 03", suggestion)

    def test_route_family(self):
        self.assertEqual(check_workbook._route_family("Fairfax 01A"), "Fairfax 01")
        self.assertEqual(check_workbook._route_family("Fairfax 01"), "Fairfax 01")
        self.assertEqual(check_workbook._route_family("Fairfax Station 12"), "Fairfax Station 12")


def write_stop_order(path: str, routes) -> None:
    import json

    with open(path, "w", encoding="utf-8") as f:
        json.dump({"routes": {r: [{"row": row} for row in rows] for r, rows in routes.items()}}, f)


class StopOrderOutputTests(unittest.TestCase):
    """Pages are printed in stop order and the checks follow it."""

    # Fairfax 07 reversed: 500 Maple Drive (row 7) first.
    ROUTES = {"Fairfax 07": [7, 3], "Fairfax 01A": [4, 9], "Fairfax 12A": [8]}

    @classmethod
    def setUpClass(cls):
        fd, cls.xlsx = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        build_season_workbook(cls.xlsx)
        fd, cls.stop_file = tempfile.mkstemp(suffix=".json")
        os.close(fd)
        write_stop_order(cls.stop_file, cls.ROUTES)
        cls.stops = route_order.load_stop_order(cls.stop_file)

    @classmethod
    def tearDownClass(cls):
        os.remove(cls.xlsx)
        os.remove(cls.stop_file)

    def test_load_stop_order(self):
        self.assertEqual(self.stops, {7: (1, 2), 3: (2, 2), 4: (1, 2), 9: (2, 2), 8: (1, 1)})

    def test_order_forms_in_stop_order(self):
        records, _ = quiet(forms.parse_order_records, self.xlsx)
        groups = forms.route_groups(records, self.stops)
        self.assertEqual([(r["address"], r["stop"]) for r in groups["Fairfax 07"]],
                         [("500 Maple Drive", "1 of 2"), ("100 Maple Drive", "2 of 2")])
        # Without a stop order: sheet order, no stop numbers.
        groups = forms.route_groups(records)
        self.assertEqual([(r["address"], r["stop"]) for r in groups["Fairfax 07"]],
                         [("100 Maple Drive", None), ("500 Maple Drive", None)])

    def test_maps_in_stop_order(self):
        group = maps.group_by_label(maps.read_records(self.xlsx))["Fairfax 07"]
        ordered = route_order.in_stop_order(group, lambda rec: rec[5], self.stops)
        self.assertEqual([rec[4] for rec in ordered], ["500 Maple Drive", "100 Maple Drive"])
        self.assertEqual([route_order.stop_label(rec[5], self.stops) for rec in ordered], ["1 of 2", "2 of 2"])

    def test_after_generating_in_stop_order(self):
        records, _ = quiet(forms.parse_order_records, self.xlsx)
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir, self.stops)
            problems, _ = quiet(check_workbook.check_after, self.xlsx, None, None, out_dir, self.stop_file)
        self.assertEqual(problems, [])

    def test_after_generating_catches_pages_not_in_stop_order(self):
        records, _ = quiet(forms.parse_order_records, self.xlsx)
        with tempfile.TemporaryDirectory() as out_dir:
            quiet(forms.save_order_forms, records, out_dir)  # sheet order
            problems, _ = quiet(check_workbook.check_after, self.xlsx, None, None, out_dir, self.stop_file)
        self.assertIn(
            "order forms: Fairfax_07.pdf page 1 should be Testerson, Erin, 500 Maple Drive (10 bags) but isn't.",
            problems,
        )

    def test_stop_order_must_cover_every_order(self):
        orders, _ = check_workbook.load_orders(self.xlsx)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "stop_order.json")
            write_stop_order(path, {"Fairfax 07": [7, 3, 5], "Fairfax 01A": [4], "Fairfax 12A": [9, 8]})
            problems = check_workbook._check_stop_order(path, orders)
        self.assertEqual(problems, [
            "stop order: 700 Birch Way (Fairfax 01A) is a stop on 'Fairfax 12A'. Rerun route_order.py.",
            "stop order: row 5 (Fairfax 07) is not an order. Rerun route_order.py.",
        ])


if __name__ == "__main__":
    unittest.main()
